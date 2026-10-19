"""
docx_builder.py
───────────────
High-level document assembler constructing clean, native editable Word documents (.docx)
from extracted and analyzed PageData, TableData, and ParagraphData structures.
"""

from __future__ import annotations
import io
import os
from typing import List, Tuple, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.pdf2word_core.models.schema import (
    PageData, TableData, ParagraphData, LineData, SpanData, ImageData, VectorCluster
)
from src.pdf2word_core.renderer.oxml_injector import OxmlInjector
from src.pdf2word_core.analyzer.style_inferencer import StyleInferencer


class DocxBuilder:
    """Assembles a fresh Word document from PageData models."""

    def __init__(self):
        self.doc = docx.Document()

    def build(self, pages: List[PageData], output_path: str) -> str:
        """Constructs and saves the .docx file."""
        if not pages:
            self.doc.save(output_path)
            return output_path

        # Configure page geometry from the first page
        first_page = pages[0]
        section = self.doc.sections[0]
        OxmlInjector.set_page_geometry(
            section=section,
            width_pt=first_page.width,
            height_pt=first_page.height,
            margins_pt=first_page.margins
        )

        for page_idx, page in enumerate(pages):
            if page_idx > 0:
                self.doc.add_page_break()

            # Render tables first if present
            for table in page.tables:
                self._render_table(table)

            # Render paragraphs
            for p in page.paragraphs:
                self._render_paragraph(p)

            # Render images
            for img in page.images:
                self._render_image(img)

        self.doc.save(output_path)
        return output_path

    def _render_table(self, table_data: TableData) -> None:
        """Renders native table with explicit fixed column widths and borders."""
        table = self.doc.add_table(rows=table_data.rows, cols=table_data.cols)
        OxmlInjector.set_table_fixed_layout(table)

        # Set column widths down to twips
        col_widths_dxa = [round(w * 20) for w in table_data.col_widths]
        OxmlInjector.set_table_grid_widths(table, col_widths_dxa)

        for cell_data in table_data.cells:
            if cell_data.row_idx >= len(table.rows) or cell_data.col_idx >= len(table.columns):
                continue
            cell = table.cell(cell_data.row_idx, cell_data.col_idx)

            # Cell borders & padding
            OxmlInjector.set_cell_borders(cell, cell_data.borders)
            OxmlInjector.set_cell_margins(cell, top_dxa=20, bottom_dxa=20, left_dxa=20, right_dxa=20)

            # Shading
            if cell_data.bg_color:
                OxmlInjector.set_cell_shading(cell, cell_data.bg_color)

            # Cell text
            cell_p = cell.paragraphs[0]
            if cell_data.lines:
                for l_idx, line in enumerate(cell_data.lines):
                    if l_idx > 0:
                        cell_p = cell.add_paragraph()
                    self._populate_paragraph(cell_p, line)
            elif cell_data.text:
                run = cell_p.add_run(cell_data.text)
                run.font.name = "Arial"
                run.font.size = Pt(9.5)

        # cantSplit on every row
        for row in table.rows:
            OxmlInjector.lock_row_anti_split(row)

    def _render_paragraph(self, p_data: ParagraphData) -> None:
        """Renders a standalone paragraph."""
        p = self.doc.add_paragraph()
        for line in p_data.lines:
            for span in line.spans:
                run = p.add_run(span.text)
                font_family = StyleInferencer.map_font_name(span.font_name)
                run.font.name = font_family
                run.font.size = Pt(span.font_size)
                run.font.bold = span.is_bold
                run.font.italic = span.is_italic
                if span.color_hex and span.color_hex != "000000":
                    try:
                        r = int(span.color_hex[0:2], 16)
                        g = int(span.color_hex[2:4], 16)
                        b = int(span.color_hex[4:6], 16)
                        run.font.color.rgb = RGBColor(r, g, b)
                    except Exception:
                        pass

        OxmlInjector.set_paragraph_spacing(
            p, before_dxa=0, after_dxa=40, line_dxa=round(p_data.line_spacing * 240)
        )

    def _populate_paragraph(self, p: docx.text.paragraph.Paragraph, line: LineData) -> None:
        """Populates paragraph with spans from a line."""
        for span in line.spans:
            run = p.add_run(span.text)
            run.font.name = StyleInferencer.map_font_name(span.font_name)
            run.font.size = Pt(max(7.0, span.font_size))
            run.font.bold = span.is_bold
            run.font.italic = span.is_italic

    def _render_image(self, img_data: ImageData) -> None:
        """Renders raster image."""
        if not img_data.image_bytes:
            return
        try:
            stream = io.BytesIO(img_data.image_bytes)
            width_in = max(0.2, min(6.5, img_data.bbox.width / 72.0))
            p = self.doc.add_paragraph()
            p.add_run().add_picture(stream, width=Inches(width_in))
        except Exception:
            pass
