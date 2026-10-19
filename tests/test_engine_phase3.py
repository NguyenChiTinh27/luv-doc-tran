"""
tests/test_engine_phase3.py
───────────────────────────
Comprehensive test suite for Phase 3 Bilingual Engine (src/engine/phase3_bilingual).

Verifies:
  - Typography constants and XML helpers (DARK_NAVY_HEX, spacing 80 dxa, keepNext, vAlign center)
  - ZeroCostGuard network interception against paid API endpoints
  - OfflineGlossary canonical GHS 16 sections, 60+ chemicals, mandatory HSE phrases, and labels
  - TokenSanitizer pre-translation spacing, unit, period, and OCR typo cleaning
  - LeakHunter post-translation ordinal, section ref, and heading repair
  - RuleBasedSynthesizer 100% offline synthesis preserving CAS numbers, formulas, and units
  - Pattern A dual-paragraph pairing (<w:keepNext/>, space_after=0, #1B365D italic space_after=4pt)
  - Pattern B single-cell line break (<w:vAlign w:val="center"/>, explicit <w:r><w:br/></w:r>)
  - ExternalLLMClient adapter configuration, offline fallback, and ZeroCostGuard integration
  - BilingualTranslator unified interface and end-to-end document pairing
  - Audit verifier verification of generated paired documents
"""

import os
import sys
import tempfile
import unittest
from unittest.mock import patch, MagicMock

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# Import Phase 3 components
from src.engine.phase3_bilingual import (
    BilingualTranslator,
    PairingEngine,
    PairingPattern,
    ZeroCostGuard,
    ApiCostViolationError,
    OfflineGlossary,
    TokenSanitizer,
    LeakHunter,
    RuleBasedSynthesizer,
    ExternalLLMClient,
    LLMProvider,
    TranslationConfig,
    BodyPair,
    CellPair,
    PairedDocument,
    DARK_NAVY_HEX,
    DARK_NAVY_RGB,
    DARK_BLUE_HEX,
    DARK_BLUE_RGB,
    BLACK_HEX,
    BLACK_RGB,
    DEFAULT_FONT_NAME,
    DEFAULT_BODY_FONT_SIZE_PT,
    DEFAULT_TABLE_FONT_SIZE_PT,
    SPACING_AFTER_DXA,
    SPACING_AFTER_PT,
    ROMAN_NUMERALS,
    add_keep_next,
    set_paragraph_spacing,
    set_cell_valign_center,
    clean_inherited_styles
)

from tests.audit_verifier import (
    verify_bilingual_docx,
    check_zero_concatenation,
    check_zero_double_tabs,
    check_pair_completeness,
    check_zero_leaked_english,
    check_header_order,
    check_visual_parity
)


class TestTypographyAndConstants(unittest.TestCase):
    """Verifies authoritative typography constants and XML styling helpers."""

    def test_typography_constants(self):
        self.assertEqual(DARK_NAVY_HEX, "1B365D")
        self.assertEqual(DARK_NAVY_RGB, (27, 54, 93))
        self.assertEqual(DARK_BLUE_HEX, "002060")
        self.assertEqual(DARK_BLUE_RGB, (0, 32, 96))
        self.assertEqual(BLACK_HEX, "000000")
        self.assertEqual(BLACK_RGB, (0, 0, 0))
        self.assertEqual(DEFAULT_FONT_NAME, "Times New Roman")
        self.assertEqual(DEFAULT_BODY_FONT_SIZE_PT, 10.5)
        self.assertEqual(DEFAULT_TABLE_FONT_SIZE_PT, 10.0)
        self.assertEqual(SPACING_AFTER_DXA, 80)
        self.assertEqual(SPACING_AFTER_PT, 4.0)
        self.assertEqual(len(ROMAN_NUMERALS), 16)
        self.assertEqual(ROMAN_NUMERALS[0], "I")
        self.assertEqual(ROMAN_NUMERALS[15], "XVI")

    def test_add_keep_next(self):
        doc = docx.Document()
        p = doc.add_paragraph("Source Line")
        add_keep_next(p)
        pPr = p._p.find(qn("w:pPr"))
        self.assertIsNotNone(pPr)
        keepNext = pPr.find(qn("w:keepNext"))
        self.assertIsNotNone(keepNext, "<w:keepNext/> must be injected into <w:pPr>")

    def test_set_paragraph_spacing(self):
        doc = docx.Document()
        p = doc.add_paragraph("Spacing Test")
        set_paragraph_spacing(p, before_dxa=0, after_dxa=80, line_dxa=240, line_rule="auto")
        pPr = p._p.find(qn("w:pPr"))
        sp = pPr.find(qn("w:spacing"))
        self.assertIsNotNone(sp)
        self.assertEqual(sp.get(qn("w:before")), "0")
        self.assertEqual(sp.get(qn("w:after")), "80")
        self.assertEqual(sp.get(qn("w:line")), "240")
        self.assertEqual(sp.get(qn("w:lineRule")), "auto")

    def test_set_cell_valign_center(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.cell(0, 0)
        set_cell_valign_center(cell)
        tcPr = cell._tc.find(qn("w:tcPr"))
        self.assertIsNotNone(tcPr)
        vAlign = tcPr.find(qn("w:vAlign"))
        self.assertIsNotNone(vAlign)
        self.assertEqual(vAlign.get(qn("w:val")), "center")

    def test_clean_inherited_styles(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        r = p.add_run("Inherited Run")
        r.italic = True
        r.font.color.rgb = RGBColor(27, 54, 93)
        self.assertTrue(len(p._p.xpath(".//w:i")) > 0)
        self.assertTrue(len(p._p.xpath(".//w:color")) > 0)

        clean_inherited_styles(p)
        self.assertEqual(len(p._p.xpath(".//w:i")), 0)
        self.assertEqual(len(p._p.xpath(".//w:color")), 0)


class TestZeroCostGuard(unittest.TestCase):
    """Verifies runtime cost protection blocking external paid APIs (R9)."""

    def test_context_manager_blocks_requests_deepseek(self):
        import requests
        with ZeroCostGuard():
            self.assertTrue(ZeroCostGuard.is_active())
            with self.assertRaises(ApiCostViolationError) as ctx:
                requests.post("https://api.deepseek.com/v1/chat/completions", json={"prompt": "hi"})
            self.assertIn("api.deepseek.com", str(ctx.exception))
            self.assertIn("R9 Cost Constraint", str(ctx.exception))
        self.assertFalse(ZeroCostGuard.is_active())

    def test_context_manager_blocks_requests_openai(self):
        import requests
        with ZeroCostGuard():
            with self.assertRaises(ApiCostViolationError) as ctx:
                requests.get("https://api.openai.com/v1/models")
            self.assertIn("api.openai.com", str(ctx.exception))

    def test_context_manager_blocks_urllib(self):
        import urllib.request
        with ZeroCostGuard():
            with self.assertRaises(ApiCostViolationError) as ctx:
                urllib.request.urlopen("https://api.deepseek.com/test")
            self.assertIn("api.deepseek.com", str(ctx.exception))

    def test_decorator_protect(self):
        import requests

        @ZeroCostGuard.protect
        def guarded_call():
            requests.post("https://api.deepseek.com/v1/translate")

        with self.assertRaises(ApiCostViolationError):
            guarded_call()

    def test_enable_and_disable_global(self):
        import requests
        ZeroCostGuard.enable()
        self.assertTrue(ZeroCostGuard.is_active())
        with self.assertRaises(ApiCostViolationError):
            requests.get("https://api.deepseek.com/status")
        ZeroCostGuard.disable()
        self.assertFalse(ZeroCostGuard.is_active())


class TestOfflineGlossary(unittest.TestCase):
    """Verifies GHS 16 sections, 60+ chemical substances, and HSE regulatory terms."""

    def test_ghs_16_sections_complete(self):
        self.assertEqual(len(OfflineGlossary.GHS_SECTIONS), 16)
        for idx in range(1, 17):
            sec = OfflineGlossary.get_ghs_section(idx)
            self.assertIsNotNone(sec)
            self.assertEqual(sec["roman"], f"{ROMAN_NUMERALS[idx-1]}.")
            self.assertTrue(len(sec["en"]) > 5)
            self.assertTrue(len(sec["vi"]) > 5)

        # Spot-check critical sections
        sec1 = OfflineGlossary.get_ghs_section(1)
        self.assertEqual(sec1["roman"], "I.")
        self.assertIn("Identification", sec1["en"])
        self.assertIn("Nhận dạng", sec1["vi"])

        sec4 = OfflineGlossary.get_ghs_section(4)
        self.assertEqual(sec4["roman"], "IV.")
        self.assertEqual(sec4["vi"], "IV. Biện pháp sơ cứu về y tế")

        sec16 = OfflineGlossary.get_ghs_section(16)
        self.assertEqual(sec16["roman"], "XVI.")
        self.assertEqual(sec16["vi"], "XVI. Thông tin khác")

    def test_chemical_substances_count_and_translations(self):
        self.assertGreaterEqual(len(OfflineGlossary.CHEMICAL_NAMES), 60)

        # Mandatory test cases from R4
        test_chems = [
            ("Silicon dioxide", "Silic đioxit"),
            ("Amorphous silicon dioxide", "Silic đioxit vô định hình"),
            ("Alumina", "Nhôm oxit"),
            ("Aluminium oxide", "Nhôm oxit"),
            ("Water", "Nước"),
            ("Calcium carbonate", "Canxi cacbonat"),
            ("Sulfuric acid", "Axit sunfuric"),
            ("Hydrochloric acid", "Axit clohydric"),
            ("Toluene", "Toluen"),
            ("Xylene", "Xylen"),
            ("Titanium dioxide", "Titan đioxit"),
            ("Tripropylene glycol diacrylate (tpgda)", "Tripropylen glycol điacrylat (TPGDA)")
        ]
        for en, expected_vi in test_chems:
            res = OfflineGlossary.lookup_chemical(en)
            self.assertEqual(res, expected_vi, f"Chemical translation mismatch for {en}")

    def test_mandatory_hse_regulatory_phrases(self):
        # R7 strict HSE terms
        hse_tests = [
            ("No smoking, drinking and eating", "Không hút thuốc, ăn uống"),
            ("Material Safety Data Sheet", "Bảng Dữ Liệu An Toàn Hóa Chất (MSDS/SDS)"),
            ("Skin sensitization", "Gây mẫn cảm da"),
            ("Eye irritation", "Gây kích ứng mắt"),
            ("Keep away from ignition sources, heat and flame", "Tránh xa các nguồn đánh lửa, nguồn nhiệt và ngọn lửa"),
            ("Get medical attention immediately", "Đưa đi cấp cứu y tế ngay lập tức")
        ]
        for en_phrase, expected_vi in hse_tests:
            res = OfflineGlossary.lookup_hse(en_phrase)
            self.assertEqual(res, expected_vi, f"HSE term mismatch for {en_phrase}")

    def test_property_labels(self):
        self.assertGreaterEqual(len(OfflineGlossary.FIELD_LABELS), 70)
        labels = [
            ("Trade name", "Tên thương phẩm"),
            ("Product name", "Tên sản phẩm"),
            ("CAS no.", "Số CAS"),
            ("Melting point", "Điểm nóng chảy"),
            ("Flash point", "Điểm chớp cháy"),
            ("Density", "Khối lượng riêng"),
            ("Viscosity", "Độ nhớt"),
            ("Solubility in water", "Độ tan trong nước")
        ]
        for en_lbl, expected_vi in labels:
            res = OfflineGlossary.lookup_label(en_lbl)
            self.assertEqual(res, expected_vi, f"Label mismatch for {en_lbl}")


class TestTokenSanitizer(unittest.TestCase):
    """Verifies pre-translation regex sanitization (R6)."""

    def test_missing_space_before_units(self):
        cases = [
            ("Wash for 15minutes with water", "Wash for 15 minutes with water"),
            ("Exposure for 2hours", "Exposure for 2 hours"),
            ("Dosage 500mg per day", "Dosage 500 mg per day"),
            ("Add 50ml solution", "Add 50 ml solution"),
            ("Limit 100ppm in air", "Limit 100 ppm in air"),
        ]
        for raw, expected in cases:
            self.assertEqual(TokenSanitizer.sanitize(raw), expected)

    def test_missing_space_after_periods(self):
        cases = [
            ("Flush eyes for 15 minutes.Assure adequate flushing.", "Flush eyes for 15 minutes. Assure adequate flushing."),
            ("Total.Consult a doctor.", "Total. Consult a doctor."),
        ]
        for raw, expected in cases:
            self.assertEqual(TokenSanitizer.sanitize(raw), expected)

    def test_ocr_typos_repair(self):
        cases = [
            ("Keep away from hest and flame", "Keep away from heat and flame"),
            ("Store away form ignition sources", "Store away from ignition sources"),
            ("Check teh label", "Check the label"),
            ("Sample recieved on Monday", "Sample received on Monday"),
            ("CITY,BINH DUONG", "CITY, BINH DUONG"),
        ]
        for raw, expected in cases:
            self.assertEqual(TokenSanitizer.sanitize(raw), expected)

    def test_camelcase_and_domain_fixes(self):
        self.assertEqual(TokenSanitizer.sanitize("Product by ScienceLab"), "Product by Science Lab")
        self.assertEqual(TokenSanitizer.sanitize("Visit ScienceLab.com"), "Visit Science Lab.com")


class TestLeakHunter(unittest.TestCase):
    """Verifies post-translation scanner repairing ordinals and headings (R6)."""

    def test_ordinals_repair(self):
        cases = [
            ("Biện pháp ở mức 1st là an toàn", "Biện pháp ở mức thứ 1 là an toàn"),
            ("Ưu tiên 2nd cần chú ý", "Ưu tiên thứ 2 cần chú ý"),
            ("Giai đoạn 3rd tiếp theo", "Giai đoạn thứ 3 tiếp theo"),
            ("Hạng mục 4th và 16th", "Hạng mục thứ 4 và thứ 16"),
        ]
        for raw, expected in cases:
            self.assertEqual(LeakHunter.hunt_and_repair(raw), expected)

    def test_section_references_repair(self):
        cases = [
            ("Xem chi tiết tại section 6.2", "Xem chi tiết tại mục 6.2"),
            ("Tuân thủ Section 4.1", "Tuân thủ Mục 4.1"),
            ("Theo section12", "Theo mục 12"),
        ]
        for raw, expected in cases:
            self.assertEqual(LeakHunter.hunt_and_repair(raw), expected)

    def test_headings_and_keywords_repair(self):
        cases = [
            ("Phần Goods này rất nguy hiểm", "Phần Hàng hóa này rất nguy hiểm"),
            ("Cần áp dụng các Measures kịp thời", "Cần áp dụng các Biện pháp kịp thời"),
            ("Trường hợp Emergency", "Trường hợp Khẩn cấp"),
            ("Xem xét Precautionary statements", "Xem xét Các cảnh báo phòng ngừa"),
        ]
        for raw, expected in cases:
            self.assertEqual(LeakHunter.hunt_and_repair(raw), expected)

    def test_detect_leaks(self):
        leaks = LeakHunter.detect_leaks("Xem mục 1st và Goods")
        self.assertTrue(len(leaks) >= 2)
        clean_text = "Xem mục thứ 1 và Hàng hóa an toàn"
        self.assertEqual(len(LeakHunter.detect_leaks(clean_text)), 0)


class TestRuleBasedSynthesizer(unittest.TestCase):
    """Verifies 100% offline synthesis, preservation of CAS, formulas, and units (R4, R7, R9)."""

    def test_ghs_section_header_normalization_to_roman(self):
        cases = [
            ("SECTION 1: Identification of the substance/mixture", "I. Nhận dạng hóa chất và thông tin về nhà cung cấp"),
            ("Section 2: Hazards Identification", "II. Nhận dạng đặc tính nguy hiểm"),
            ("SECTION 3: Composition / Information on Ingredients", "III. Thông tin về thành phần các chất"),
            ("Section 4: First Aid Measures", "IV. Biện pháp sơ cứu về y tế"),
            ("SECTION 9: Physical and Chemical Properties", "IX. Đặc tính lý, hóa học"),
            ("SECTION 16: Other Information", "XVI. Thông tin khác"),
        ]
        for en, expected_vi in cases:
            vi = RuleBasedSynthesizer.synthesize(en)
            self.assertEqual(vi, expected_vi)

    def test_cas_numbers_preservation(self):
        cas_numbers = ["7631-86-9", "1344-28-1", "7732-18-5", "471-34-1", "108-88-3"]
        for cas in cas_numbers:
            # Standalone CAS
            self.assertEqual(RuleBasedSynthesizer.synthesize(cas), cas)
            self.assertTrue(RuleBasedSynthesizer.is_pure_data(cas))
            # CAS with label
            result = RuleBasedSynthesizer.synthesize(f"CAS no.: {cas}")
            self.assertIn(cas, result)
            self.assertIn("Số CAS", result)

    def test_chemical_formulas_preservation(self):
        formulas = ["CaCO3", "SiO2", "Al2O3", "H2SO4", "H2O", "NaOH", "NaCl"]
        for form in formulas:
            self.assertTrue(RuleBasedSynthesizer.is_pure_data(form))
            res = RuleBasedSynthesizer.synthesize(f"Chemical formula: {form}")
            self.assertIn(form, res)
            self.assertIn("Công thức hóa học", res)

    def test_numeric_ranges_and_units_preservation(self):
        cases = [
            (": 1.5 (±0.05)", ": 1.5 (±0.05)"),
            ("95-100%", "95-100%"),
            ("> 1700 °C", "> 1700 °C"),
            ("15 minutes", "15 phút"),
            ("pH: 7.0 - 8.5", "Giá trị pH: 7.0 - 8.5"),
        ]
        for en, expected_pattern in cases:
            res = RuleBasedSynthesizer.synthesize(en)
            self.assertIn(expected_pattern.split()[-1], res)

    def test_label_value_deconstruction(self):
        res = RuleBasedSynthesizer.synthesize("Trade name: Polyurethane Resin")
        self.assertTrue(res.startswith("Tên thương phẩm:"))
        self.assertIn("Polyuretan", res)

    def test_hse_phrase_synthesis(self):
        res = RuleBasedSynthesizer.synthesize("No smoking, drinking and eating.")
        self.assertEqual(res, "Không hút thuốc, ăn uống.")


class TestPairingEngine(unittest.TestCase):
    """Verifies Pattern A dual-paragraph and Pattern B table cell typography (R1, R8)."""

    def test_pattern_a_dual_paragraph(self):
        doc = docx.Document()
        p_src, p_trans_el = PairingEngine.apply_pattern_a(
            doc,
            en_text="SECTION 1: Identification",
            vi_text="I. Nhận dạng hóa chất và thông tin về nhà cung cấp",
            is_bold=True
        )

        # 1. Source paragraph must have <w:keepNext/>
        pPr_src = p_src._p.find(qn("w:pPr"))
        self.assertIsNotNone(pPr_src)
        self.assertIsNotNone(pPr_src.find(qn("w:keepNext")), "Source line must have <w:keepNext/>")

        # 2. Source paragraph must have space_after=0
        sp_src = pPr_src.find(qn("w:spacing"))
        self.assertIsNotNone(sp_src)
        self.assertEqual(sp_src.get(qn("w:after")), "0")

        # 3. Source paragraph color must be black and not italic
        self.assertEqual(p_src.runs[0].font.color.rgb, BLACK_RGB)
        self.assertFalse(bool(p_src.runs[0].font.italic))

        # 4. Translated paragraph element must immediately succeed source paragraph
        self.assertIs(p_src._p.getnext(), p_trans_el)

        # 5. Translated paragraph must have space_after=80 dxa (4 pt)
        pPr_trans = p_trans_el.find(qn("w:pPr"))
        sp_trans = pPr_trans.find(qn("w:spacing"))
        self.assertEqual(sp_trans.get(qn("w:after")), "80")

        # 6. Translated run must have italic <w:i/> and color #1B365D
        r_trans = p_trans_el.find(qn("w:r"))
        rPr_trans = r_trans.find(qn("w:rPr"))
        self.assertIsNotNone(rPr_trans.find(qn("w:i")), "Translation must be italic")
        col = rPr_trans.find(qn("w:color"))
        self.assertIsNotNone(col)
        self.assertEqual(col.get(qn("w:val")), "1B365D", "Translation must use Dark Navy #1B365D")

    def test_pattern_b_single_cell(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.cell(0, 0)

        PairingEngine.apply_pattern_b(
            cell,
            en_text="Silicon dioxide",
            vi_text="Silic đioxit",
            is_bold=False
        )

        # 1. Cell vertical alignment must be center
        tcPr = cell._tc.find(qn("w:tcPr"))
        vAlign = tcPr.find(qn("w:vAlign"))
        self.assertIsNotNone(vAlign)
        self.assertEqual(vAlign.get(qn("w:val")), "center")

        # 2. Must contain exactly one paragraph
        self.assertEqual(len(cell.paragraphs), 1)
        p = cell.paragraphs[0]

        # 3. Must contain explicit isolated line break run <w:r><w:br/></w:r>
        br_elems = p._p.xpath(".//w:r/w:br")
        self.assertEqual(len(br_elems), 1, "Pattern B cell must have explicit <w:br/>")

        # 4. Translated run must be italic and Dark Navy
        vi_run = p.runs[-1]
        self.assertTrue(vi_run.font.italic)
        self.assertEqual(vi_run.font.color.rgb, RGBColor(0x1B, 0x36, 0x5D))

    def test_pair_document_full_flow(self):
        doc = docx.Document()
        doc.add_paragraph("SECTION 1: Identification of the substance/mixture and of the company/undertaking")
        doc.add_paragraph("Trade name: Alumina Solution")

        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Ingredient"
        tbl.cell(0, 1).text = "CAS no."
        tbl.cell(1, 0).text = "Water"
        tbl.cell(1, 1).text = "7732-18-5"

        paired_doc = PairingEngine.pair_document(doc)
        self.assertGreaterEqual(len(paired_doc.body_pairs), 2)
        # Note: 7732-18-5 is pure CAS number so it should not be translated
        self.assertGreaterEqual(len(paired_doc.cell_pairs), 2)


class TestExternalLLMClient(unittest.TestCase):
    """Verifies optional external LLM adapter configuration and safety guard."""

    def test_unconfigured_client_behavior(self):
        client = ExternalLLMClient(api_key="", fallback_to_offline=True)
        self.assertFalse(client.is_configured())
        # Falls back to offline synthesizer without crashing
        res = client.translate("Water")
        self.assertEqual(res, "Nước")

        # When fallback is disabled and no key, raises ValueError
        strict_client = ExternalLLMClient(api_key="", fallback_to_offline=False)
        with self.assertRaises(ValueError):
            strict_client.translate("Water")

    def test_zero_cost_guard_blocks_client_network_call(self):
        client = ExternalLLMClient(
            api_key="test-dummy-key-12345",
            base_url="https://api.deepseek.com/v1",
            fallback_to_offline=False
        )
        self.assertTrue(client.is_configured())

        with ZeroCostGuard():
            with self.assertRaises(ApiCostViolationError):
                client.translate("Hazardous chemical")


class TestBilingualTranslatorUnified(unittest.TestCase):
    """Verifies the unified BilingualTranslator facade."""

    def test_default_offline_initialization(self):
        translator = BilingualTranslator()
        self.assertEqual(translator.mode, "offline")
        self.assertTrue(translator.zero_cost)
        self.assertTrue(ZeroCostGuard.is_active())

        # Direct translation
        res = translator.translate_text("Sulfuric acid")
        self.assertEqual(res, "Axit sunfuric")

    def test_pair_paragraph_and_cell(self):
        translator = BilingualTranslator()
        doc = docx.Document()
        p_src, p_trans = translator.pair_paragraph(doc, "SECTION 4: First Aid Measures")
        self.assertIsNotNone(p_src)
        self.assertIsNotNone(p_trans)

        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.cell(0, 0)
        translator.pair_cell(cell, "Calcium carbonate")
        self.assertIn("Canxi cacbonat", cell.text)

    def test_process_docx_file(self):
        translator = BilingualTranslator()
        with tempfile.TemporaryDirectory() as tmpdir:
            input_docx = os.path.join(tmpdir, "test_input.docx")
            output_docx = os.path.join(tmpdir, "test_output.docx")

            doc = docx.Document()
            doc.add_paragraph("SECTION 1: Identification")
            doc.add_paragraph("Product name: Toluene")
            doc.save(input_docx)

            res_path = translator.process_docx(input_docx, output_docx)
            self.assertTrue(os.path.exists(res_path))

            # Verify generated output
            out_doc = docx.Document(res_path)
            full_text = " ".join([p.text for p in out_doc.paragraphs])
            self.assertIn("Nhận dạng", full_text)
            self.assertIn("Toluen", full_text)


class TestAuditVerifierIntegration(unittest.TestCase):
    """Verifies that documents created by Phase 3 pass the authoritative 7-Point Audit Verifier."""

    def test_bilingual_document_passes_audit_verifier(self):
        translator = BilingualTranslator()
        doc = docx.Document()

        # Section 1
        translator.pair_paragraph(doc, "SECTION 1: Identification of the substance/mixture and of the company/undertaking", is_bold=True)
        translator.pair_paragraph(doc, "Trade name: Titanium dioxide coating")
        translator.pair_paragraph(doc, "Product code: TiO2-99")

        # Section 2
        translator.pair_paragraph(doc, "SECTION 2: Hazards Identification", is_bold=True)
        translator.pair_paragraph(doc, "Causes skin irritation.")
        translator.pair_paragraph(doc, "Causes serious eye irritation.")

        # Section 3 with table
        translator.pair_paragraph(doc, "SECTION 3: Composition / Information on Ingredients", is_bold=True)
        tbl = doc.add_table(rows=3, cols=3)
        headers = ["Chemical Name", "CAS no.", "Weight %"]
        for j, h in enumerate(headers):
            translator.pair_cell(tbl.cell(0, j), h, is_bold=True)

        # Row 1
        translator.pair_cell(tbl.cell(1, 0), "Titanium dioxide")
        tbl.cell(1, 1).text = "13463-67-7"
        tbl.cell(1, 2).text = "60-70%"

        # Row 2
        translator.pair_cell(tbl.cell(2, 0), "Water")
        tbl.cell(2, 1).text = "7732-18-5"
        tbl.cell(2, 2).text = "30-40%"

        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "audit_test.docx")
            doc.save(out_path)

            audit_res = verify_bilingual_docx(out_path)
            self.assertEqual(audit_res["concatenation_violations"], 0, "Check 1: Concat violations must be 0")
            self.assertEqual(audit_res["double_tab_violations"], 0, "Check 2: Double tab violations must be 0")
            self.assertEqual(audit_res["pair_completeness"], 1.0, "Check 3: Pair completeness must be 1.0 (100%)")
            self.assertEqual(audit_res["leaked_english_count"], 0, "Check 4: Leaked english must be 0")
            self.assertEqual(audit_res["header_order_violations"], 0, "Check 5: Header order violations must be 0")
            self.assertEqual(audit_res["clipped_row_violations"], 0, "Check 6: Clipped row violations must be 0")


if __name__ == "__main__":
    unittest.main()
