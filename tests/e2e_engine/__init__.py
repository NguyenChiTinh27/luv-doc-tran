"""
tests/e2e_engine/__init__.py
─────────────────────────────────────────────────────────────────────────────
End-to-End Opaque-Box Test Harness for PDF-to-Bilingual-Word Conversion & QC Pipeline.

Provides unified fixture helpers, import adapters, and diagnostic verifiers
supporting Tiers 1 through 4:
  - Tier 1: Feature Coverage (F1 to F15)
  - Tier 2: Boundary Value Analysis & Corner Cases
  - Tier 3: Pairwise Cross-Feature Interactions
  - Tier 4: Real-World Scenarios & Golden Audit Verification
─────────────────────────────────────────────────────────────────────────────
"""

import os
import sys
import re
import tempfile
import zipfile
import lxml.etree as etree
from typing import Dict, Any, List, Optional, Tuple

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, RGBColor

# Import authoritative 7-point audit verifier primitives
from tests.audit_verifier import (
    verify_bilingual_docx,
    check_zero_concatenation,
    check_zero_double_tabs,
    check_pair_completeness,
    check_zero_leaked_english,
    check_header_order,
    check_zero_clipped_rows,
    check_visual_parity,
    is_translatable_source_text,
    is_vietnamese_text,
    is_chemical_formula,
    WHITELIST_TOKENS,
    NS,
    W_NS
)

# Typography Constants
DARK_NAVY_HEX = "1B365D"
DARK_NAVY_RGB = RGBColor(0x1B, 0x36, 0x5D)
BLACK_RGB = RGBColor(0, 0, 0)
DEFAULT_FONT_NAME = "Times New Roman"

# Attempt to resolve from src.engine, with reference fallback
try:
    from src.engine.phase3_bilingual.cost_guard import ZeroCostGuard, ApiCostViolationError
except ImportError:
    try:
        from core.bilingual_qc_v2 import ZeroCostGuard, ApiCostViolationError
    except ImportError:
        class ApiCostViolationError(Exception):
            pass

        class ZeroCostGuard:
            def __enter__(self):
                return self
            def __exit__(self, exc_type, exc_val, exc_tb):
                pass
            @classmethod
            def protect(cls, fn):
                return fn


def create_empty_docx(path: str) -> str:
    """Create a minimal valid Word document at the given path."""
    doc = docx.Document()
    doc.save(path)
    return path


def apply_pattern_a(doc: docx.Document, en_text: str, vi_text: str) -> Tuple[docx.text.paragraph.Paragraph, docx.text.paragraph.Paragraph]:
    """
    Applies Pattern A (Dual-Paragraph Pair for Body/Lists):
    - English paragraph: <w:keepNext/>, space_after = 0pt
    - Vietnamese paragraph: Dark Navy #1B365D, italic, space_before = 0pt, space_after = 4pt (80 dxa)
    """
    p_en = doc.add_paragraph()
    p_en.paragraph_format.space_before = Pt(0)
    p_en.paragraph_format.space_after = Pt(0)
    p_en.paragraph_format.keep_with_next = True
    r_en = p_en.add_run(en_text)
    r_en.font.name = DEFAULT_FONT_NAME
    r_en.font.color.rgb = BLACK_RGB

    p_vi = doc.add_paragraph()
    p_vi.paragraph_format.space_before = Pt(0)
    p_vi.paragraph_format.space_after = Pt(4)
    r_vi = p_vi.add_run(vi_text)
    r_vi.font.name = DEFAULT_FONT_NAME
    r_vi.font.italic = True
    r_vi.font.color.rgb = DARK_NAVY_RGB

    return p_en, p_vi


def apply_pattern_b(cell: docx.table._Cell, en_text: str, vi_text: str) -> docx.text.paragraph.Paragraph:
    """
    Applies Pattern B (Single-Cell Line Break for Table Cells):
    - Table cell: <w:vAlign w:val="center"/>
    - English run: Black, normal
    - Separator: explicit <w:r><w:br/></w:r>
    - Vietnamese run: Dark Navy #1B365D, italic
    """
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)

    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    r_en = p.add_run(en_text)
    r_en.font.name = DEFAULT_FONT_NAME
    r_en.font.color.rgb = BLACK_RGB

    # Explicit break run
    r_br = p.add_run()
    r_br.add_break()

    r_vi = p.add_run(vi_text)
    r_vi.font.name = DEFAULT_FONT_NAME
    r_vi.font.italic = True
    r_vi.font.color.rgb = DARK_NAVY_RGB

    return p


def unlock_table_row_heights(table: docx.table.Table) -> None:
    """Unlocks row heights from exact to atLeast and injects cantSplit."""
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        for child in list(trPr):
            if child.tag.endswith("trHeight"):
                trPr.remove(child)
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "240")
        trHeight.set(qn("w:hRule"), "atLeast")
        trPr.append(trHeight)

        if not any(c.tag.endswith("cantSplit") for c in trPr):
            trPr.append(OxmlElement("w:cantSplit"))


def promote_native_header_footer(doc: docx.Document, title: str = "Bilingual MSDS Document") -> None:
    """Configures native Word section header with dynamic PAGE / NUMPAGES fields."""
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.text = f"{title} | "
    
    # Dynamic PAGE field
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    
    r = p.add_run(" of ")
    
    # Dynamic NUMPAGES field
    fld_num = OxmlElement("w:fldSimple")
    fld_num.set(qn("w:instr"), "NUMPAGES")
    p._p.append(fld_num)


# =============================================================================
# UNIFIED AUDIT HELPER FUNCTIONS FOR TESTS
# =============================================================================

def audit_docx(docx_path: str) -> Dict[str, Any]:
    """Execute complete 7-Point OpenXML Audit on target docx."""
    return verify_bilingual_docx(docx_path)


def audit_check1_concatenation(docx_path: str) -> Tuple[int, List[Dict[str, Any]]]:
    """Inspects Check 1 Zero Concatenation across all XML trees in docx."""
    with zipfile.ZipFile(docx_path, "r") as z:
        trees = []
        for name in z.namelist():
            if name == "word/document.xml" or re.match(r"word/(header|footer)\d+\.xml", name):
                try:
                    trees.append(etree.fromstring(z.read(name)))
                except Exception:
                    pass
        return check_zero_concatenation(trees)


def audit_check2_double_tabs(docx_path: str) -> List[str]:
    """Inspects Check 2 Zero Double-Tabs."""
    with zipfile.ZipFile(docx_path, "r") as z:
        tree = etree.fromstring(z.read("word/document.xml"))
        count, violations = check_zero_double_tabs(tree)
        return violations


def audit_check3_pair_completeness(docx_path: str) -> Dict[str, Any]:
    """Inspects Check 3 Pair Completeness."""
    with zipfile.ZipFile(docx_path, "r") as z:
        tree = etree.fromstring(z.read("word/document.xml"))
        completeness, total_src, validated, unpaired = check_pair_completeness(tree)
        return {
            "ratio": completeness,
            "total_source_blocks": total_src,
            "validated_pairs": validated,
            "unpaired_examples": unpaired
        }


def audit_check4_leaked_english(docx_path: str) -> List[Dict[str, Any]]:
    """Inspects Check 4 Zero Leaked English."""
    with zipfile.ZipFile(docx_path, "r") as z:
        tree = etree.fromstring(z.read("word/document.xml"))
        count, violations = check_zero_leaked_english(tree)
        return violations


def audit_check5_header_order(docx_path: str) -> List[Dict[str, Any]]:
    """Inspects Check 5 Strict Header Ordering."""
    with zipfile.ZipFile(docx_path, "r") as z:
        tree = etree.fromstring(z.read("word/document.xml"))
        count, violations = check_header_order(tree)
        return violations


def audit_check6_clipped_rows(docx_path: str) -> List[str]:
    """Inspects Check 6 Zero Clipped Rows."""
    with zipfile.ZipFile(docx_path, "r") as z:
        tree = etree.fromstring(z.read("word/document.xml"))
        count, violations = check_zero_clipped_rows(tree)
        return violations


def audit_check7_visual_parity(docx_path: str) -> Dict[str, Any]:
    """Inspects Check 7 Visual Parity & Header/Footer Promotion."""
    with zipfile.ZipFile(docx_path, "r") as z:
        tree = etree.fromstring(z.read("word/document.xml"))
        passed, issues = check_visual_parity(z, tree)
        return {"passed": passed, "issues": issues}
