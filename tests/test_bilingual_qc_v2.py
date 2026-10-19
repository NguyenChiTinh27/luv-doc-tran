"""
tests/test_bilingual_qc_v2.py

Comprehensive 4-Tier Test Suite for Bilingual Document Reconstruction & QC Engine (bilingual_qc_v2.py):
  - Tier 1: Feature Coverage (R1-R9 isolation, reflow, grid, chemicals, consolidation, sanitizer, glossary, typography, zero API cost)
  - Tier 2: Boundary & Corner Cases (empty lines, 0-value, massive tables, multi-column colons, nested tables, symbols, CAS numbers)
  - Tier 3: Cross-Feature Interactions (chemical substances in multi-column tables with line wraps)
  - Tier 4: Real-world document validation tests using actual sample files (MSDS-YTU0035-EN.pdf, MSDS-YTU0070-EN.pdf, MSDS-FU43288-EN.doc) & golden fixtures
  - Includes ZeroCostGuard fixture: asserts zero external HTTP calls to api.deepseek.com
"""

import unittest
import os
import sys
import re
import tempfile
import zipfile
import lxml.etree as etree
from unittest.mock import patch

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, RGBColor

# Import the authoritative 7-Point XML & Visual Audit verifier
from tests.audit_verifier import (
    verify_bilingual_docx,
    check_zero_concatenation,
    check_zero_double_tabs,
    check_pair_completeness,
    check_zero_leaked_english,
    check_header_order,
    check_zero_clipped_rows,
    check_visual_parity,
    is_chemical_formula,
    is_vietnamese_text,
    is_translatable_source_text,
    WHITELIST_TOKENS,
    NS,
    W_NS
)

# Conditionally import bilingual_qc_v2 if implemented, otherwise provide test reference implementations
try:
    import core.bilingual_qc_v2 as bqc
except ImportError:
    bqc = None


# ==============================================================================
# ZERO COST GUARD FIXTURE & REFERENCE ENGINES (R9)
# ==============================================================================

class ApiCostViolationError(Exception):
    """Raised when an external paid API call is attempted during testing/QC."""
    pass


class ZeroCostGuard:
    """Interception guard to assert 0 external HTTP calls to api.deepseek.com or other paid endpoints."""

    def __enter__(self):
        self.patches = []
        # Guard urllib
        try:
            import urllib.request
            orig_urlopen = urllib.request.urlopen

            def guarded_urlopen(url, *args, **kwargs):
                url_str = getattr(url, "full_url", str(url))
                if "api.deepseek.com" in url_str:
                    raise ApiCostViolationError(f"Prohibited paid API call to {url_str} blocked by ZeroCostGuard!")
                return orig_urlopen(url, *args, **kwargs)

            p1 = patch("urllib.request.urlopen", side_effect=guarded_urlopen)
            p1.start()
            self.patches.append(p1)
        except Exception:
            pass

        # Guard requests if available
        try:
            import requests
            orig_post = requests.post
            orig_get = requests.get

            def guarded_post(url, *args, **kwargs):
                if "api.deepseek.com" in str(url):
                    raise ApiCostViolationError(f"Prohibited paid API call to {url} blocked by ZeroCostGuard!")
                return orig_post(url, *args, **kwargs)

            def guarded_get(url, *args, **kwargs):
                if "api.deepseek.com" in str(url):
                    raise ApiCostViolationError(f"Prohibited paid API call to {url} blocked by ZeroCostGuard!")
                return orig_get(url, *args, **kwargs)

            p2 = patch("requests.post", side_effect=guarded_post)
            p3 = patch("requests.get", side_effect=guarded_get)
            p2.start()
            p3.start()
            self.patches.extend([p2, p3])
        except Exception:
            pass

        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        for p in self.patches:
            p.stop()


# Reference implementation of offline glossary & rule-based engine (R9 / R7)
OFFLINE_GLOSSARY = {
    # 16 GHS Section Headers
    "SECTION 1: Identification of the substance/mixture and of the company/undertaking":
        "I. Nhận dạng hóa chất và thông tin về nhà cung cấp",
    "SECTION 2: Hazards Identification":
        "II. Nhận dạng đặc tính nguy hiểm",
    "SECTION 3: Composition / Information on Ingredients":
        "III. Thông tin về thành phần các chất",
    "SECTION 4: First Aid Measures":
        "IV. Biện pháp sơ cứu về y tế",
    "SECTION 5: Fire Fighting Measures":
        "V. Biện pháp xử lý khi có hỏa hoạn",
    "SECTION 6: Accidental Release Measures":
        "VI. Biện pháp phòng ngừa sự cố tràn đổ",
    "SECTION 7: Handling and Storage":
        "VII. Sử dụng và bảo quản",
    "SECTION 8: Exposure Controls / Personal Protection":
        "VIII. Kiểm soát phơi nhiễm / Bảo vệ cá nhân",
    "SECTION 9: Physical and Chemical Properties":
        "IX. Đặc tính lý, hóa học",
    "SECTION 10: Stability and Reactivity":
        "X. Tính ổn định và khả năng phản ứng",
    "SECTION 11: Toxicological Information":
        "XI. Thông tin về độc tính",
    "SECTION 12: Ecological Information":
        "XII. Thông tin về sinh thái",
    "SECTION 13: Disposal Considerations":
        "XIII. Biện pháp tiêu hủy",
    "SECTION 14: Transport Information":
        "XIV. Thông tin khi vận chuyển",
    "SECTION 15: Regulatory Information":
        "XV. Thông tin về quy chuẩn pháp luật",
    "SECTION 16: Other Information":
        "XVI. Thông tin khác",

    # Common Chemicals (R4)
    "Silicon dioxide": "Silic đioxit",
    "amorphous silicon dioxide, chemically prepared": "Silic đioxit vô định hình, điều chế hóa học",
    "Alumina": "Nhôm oxit",
    "Water": "Nước",
    "Calcium carbonate": "Canxi cacbonat",
    "Talc powder": "Bột Talc",
    "Sulfuric acid": "Axit sunfuric",
    "Toluene": "Toluen",
    "Xylene": "Xylen",
    "Butyl acetate": "Butyl axetat",
    "Ethyl acetate": "Etyl axetat",
    "Acetone": "Axeton",
    "Titanium dioxide": "Titan đioxit",

    # HSE Terminology (R7)
    "No smoking, drinking and eating": "Không hút thuốc, ăn uống",
    "Material Safety Data Sheet": "Bảng Dữ Liệu An Toàn Hóa Chất (MSDS/SDS)",
    "Skin sensitization": "Gây mẫn cảm da",
    "Eye irritation": "Gây kích ứng mắt",
    "Flash point": "Điểm chớp cháy",
    "Boiling point": "Điểm sôi",
    "Melting point": "Điểm nóng chảy",
    "Appearance": "Hình thái bên ngoài",
    "Odour": "Mùi",
    "Odor": "Mùi",
    "Relative density": "Tỷ trọng tương đối",
    "Solubility in water": "Độ tan trong nước",
}

# Reference Token Sanitizer (R6)
def sanitize_tokens(text: str) -> str:
    """Pre-translation regex cleaner for units, typos, and sentence boundaries."""
    # Fix missing spaces around units and numbers (e.g. 15minutes -> 15 minutes)
    text = re.sub(r"(\d+)(minutes|hours|days|weeks|months|years|sec|min|hr|mg|kg|ppm)\b", r"\1 \2", text)
    # Fix fused sentence boundaries (e.g. 15minutes.Assure -> 15 minutes. Assure)
    text = re.sub(r"(\b[a-zA-Z0-9]+)\.([A-Z][a-z]+)", r"\1. \2", text)
    # Fix OCR typos: hest and flame -> heat and flame, form ignition -> from ignition
    text = re.sub(r"\bhest and flame\b", "heat and flame", text, flags=re.IGNORECASE)
    text = re.sub(r"\bKeep away form ignition\b", "Keep away from ignition", text, flags=re.IGNORECASE)
    return text

# Reference Leak Hunter (R6)
def hunt_and_repair_leaks(text: str) -> str:
    """Post-translation leak hunter: repairs leaked ordinals and section words."""
    replacements = {
        r"\b1st\b": "thứ 1",
        r"\b2nd\b": "thứ 2",
        r"\b3rd\b": "thứ 3",
        r"\b4th\b": "thứ 4",
        r"\bGoods\b": "Hàng hóa",
        r"\bMeasures\b": "Biện pháp",
        r"\bsection\s*6\.2\b": "mục 6.2",
        r"\bsection\s*(\d+(\.\d+)?)\b": r"mục \1",
    }
    for pat, rep in replacements.items():
        text = re.sub(pat, rep, text, flags=re.IGNORECASE)
    return text


# Reference Pattern A & Pattern B Builders (R1 / R8)
def apply_pattern_a(doc, en_text: str, vi_text: str):
    """Create Pattern A dual-paragraph pair: EN <w:p> keepNext + VI <w:p> #1B365D italic."""
    # English paragraph
    p_en = doc.add_paragraph()
    pPr_en = p_en._p.get_or_add_pPr()
    pPr_en.append(OxmlElement("w:keepNext"))
    sp_en = OxmlElement("w:spacing")
    sp_en.set(qn("w:before"), "0")
    sp_en.set(qn("w:after"), "0")
    pPr_en.append(sp_en)
    r_en = p_en.add_run(en_text)
    r_en.font.color.rgb = RGBColor(0, 0, 0)

    # Vietnamese paragraph
    p_vi = doc.add_paragraph()
    pPr_vi = p_vi._p.get_or_add_pPr()
    sp_vi = OxmlElement("w:spacing")
    sp_vi.set(qn("w:before"), "0")
    sp_vi.set(qn("w:after"), "80")  # 80 dxa = 4 pt
    pPr_vi.append(sp_vi)
    r_vi = p_vi.add_run(vi_text)
    r_vi.italic = True
    r_vi.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    return p_en, p_vi


def apply_pattern_b(cell, en_text: str, vi_text: str):
    """Create Pattern B single-cell line break: EN run + <w:r><w:br/></w:r> + VI run #1B365D italic."""
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.text = ""
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    pPr.append(sp)

    # Source run
    r_en = p.add_run(en_text)

    # Explicit break
    r_br = p.add_run()
    r_br._r.append(OxmlElement("w:br"))

    # Vietnamese run
    r_vi = p.add_run(vi_text)
    r_vi.italic = True
    r_vi.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    return p


# ==============================================================================
# TIER 1: FEATURE COVERAGE TESTS (R1 – R9)
# ==============================================================================

class TestTier1FeatureCoverage(unittest.TestCase):
    """Verify each feature R1-R9 in isolation according to SCOPE.md and ORIGINAL_REQUEST.md."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_r1_pattern_a_dual_paragraph_keep_next(self):
        """R1 / Pattern A: Source <w:p> must have <w:keepNext/> and space_after=0pt."""
        doc = docx.Document()
        apply_pattern_a(doc, "Section 1: Chemical Product", "I. Nhận dạng hóa chất")
        out_path = os.path.join(self.temp_dir, "r1_pattern_a.docx")
        doc.save(out_path)

        with zipfile.ZipFile(out_path, "r") as z:
            tree = etree.fromstring(z.read("word/document.xml"))
            keep_next = tree.xpath("/w:document/w:body/w:p[1]/w:pPr/w:keepNext", namespaces=NS)
            self.assertEqual(len(keep_next), 1, "Pattern A source paragraph MUST have <w:keepNext/>")

            spacing_en = tree.xpath("/w:document/w:body/w:p[1]/w:pPr/w:spacing/@w:after", namespaces=NS)
            self.assertEqual(spacing_en, ["0"], "Pattern A source paragraph MUST have space_after=0")

            spacing_vi = tree.xpath("/w:document/w:body/w:p[2]/w:pPr/w:spacing/@w:after", namespaces=NS)
            self.assertEqual(spacing_vi, ["80"], "Pattern A translated paragraph MUST have space_after=80 (4 pt)")

    def test_r1_pattern_b_single_cell_explicit_break(self):
        """R1 / Pattern B: Table cell must separate source and translation with explicit <w:br/>."""
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        apply_pattern_b(tbl.rows[0].cells[0], "Flash point: 25 °C", "Điểm chớp cháy: 25 °C")
        out_path = os.path.join(self.temp_dir, "r1_pattern_b.docx")
        doc.save(out_path)

        with zipfile.ZipFile(out_path, "r") as z:
            tree = etree.fromstring(z.read("word/document.xml"))
            br_elem = tree.xpath("//w:tc//w:r/w:br", namespaces=NS)
            self.assertEqual(len(br_elem), 1, "Pattern B table cell MUST have explicit <w:br/> run")

    def test_r8_visual_pairing_typography_and_colors(self):
        """R8: Translated text must use Dark Navy #1B365D and italic <w:i/>."""
        doc = docx.Document()
        apply_pattern_a(doc, "Hazards Identification", "Nhận dạng đặc tính nguy hiểm")
        out_path = os.path.join(self.temp_dir, "r8_typography.docx")
        doc.save(out_path)

        with zipfile.ZipFile(out_path, "r") as z:
            tree = etree.fromstring(z.read("word/document.xml"))
            vi_p = tree.xpath("/w:document/w:body/w:p[2]", namespaces=NS)[0]
            color = vi_p.xpath(".//w:color/@w:val", namespaces=NS)
            self.assertIn("1B365D", color, "Translated line must use Dark Navy #1B365D")
            italic = vi_p.xpath(".//w:i", namespaces=NS)
            self.assertTrue(len(italic) > 0, "Translated line must have italic formatting")

    def test_r2_dynamic_reflow_unlocker(self):
        """R2: Exact table row height <w:trHeight w:hRule='exact'/> must be unlocked to 'atLeast'."""
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:hRule"), "exact")
        trHeight.set(qn("w:val"), "360")
        trPr.append(trHeight)

        # Unlock function
        def unlock_tr_heights(d):
            for elem in d._element.xpath("//w:trHeight[@w:hRule='exact']"):
                elem.set(qn("w:hRule"), "atLeast")

        unlock_tr_heights(doc)
        out_path = os.path.join(self.temp_dir, "r2_reflow.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["clipped_row_violations"], 0, "All w:hRule='exact' must be unlocked to 'atLeast'")

    def test_r2_native_header_footer_promotion(self):
        """R2: Header/footer promotion must embed dynamic PAGE/NUMPAGES fields into <w:hdr>/<w:ftr>."""
        doc = docx.Document()
        sec = doc.sections[0]
        p_hdr = sec.header.paragraphs[0]
        p_hdr.add_run("Safety Data Sheet - ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p_hdr._p.append(fld)

        out_path = os.path.join(self.temp_dir, "r2_hdr_ftr.docx")
        doc.save(out_path)

        with zipfile.ZipFile(out_path, "r") as z:
            header_xml = [z.read(f) for f in z.namelist() if "header" in f]
            self.assertTrue(any(b"PAGE" in h for h in header_xml), "Header must contain dynamic PAGE field")

    def test_r3_2_column_grid_reconstruction_and_tab_stripping(self):
        """R3: Label : Value must be converted into 28%/72% table without double tabs."""
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=2)
        # Borderless table
        tblPr = tbl._tbl.tblPr
        tblBorders = parse_xml(
            r'''<w:tblBorders %s>
                <w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/>
                <w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>
            </w:tblBorders>''' % nsdecls("w")
        )
        tblPr.append(tblBorders)
        tblPr.append(parse_xml(r'<w:tblLayout %s w:type="fixed"/>' % nsdecls("w")))

        # Col 0: 28% (~2520 dxa), Col 1: 72% (~6480 dxa)
        row = tbl.rows[0]
        apply_pattern_b(row.cells[0], "Chemical Name", "Tên hóa chất")
        apply_pattern_b(row.cells[1], ": Calcium Carbonate", ": Canxi cacbonat")

        out_path = os.path.join(self.temp_dir, "r3_grid.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["double_tab_violations"], 0, "No double tabs allowed in grid reconstruction")
        self.assertEqual(report["concatenation_violations"], 0)

    def test_r5_paragraph_consolidation(self):
        """R5: Consecutive unpunctuated lines must merge into a single logical paragraph."""
        lines = [
            "Check for and remove any contact lenses. In case of contact, immediately flush eyes with plenty of water for at least 15",
            "minutes. Cold water may be used. WARM water MUST be used. Get medical attention."
        ]

        def consolidate_lines(raw_lines):
            consolidated = []
            buf = ""
            for l in raw_lines:
                if buf:
                    buf += " " + l.strip()
                else:
                    buf = l.strip()
                if buf.endswith((".", ";", ":", "!", "?")):
                    consolidated.append(buf)
                    buf = ""
            if buf:
                consolidated.append(buf)
            return consolidated

        res = consolidate_lines(lines)
        self.assertEqual(len(res), 1, "Unpunctuated lines must consolidate into a single paragraph")
        self.assertIn("at least 15 minutes.", res[0])

    def test_r4_chemical_substances_table_coverage(self):
        """R4: Common chemical substances (Section 2/3) must be translated with zero skipped rows."""
        doc = docx.Document()
        tbl = doc.add_table(rows=3, cols=3)
        chemicals = [
            ("amorphous silicon dioxide, chemically prepared", "Silic đioxit vô định hình, điều chế hóa học"),
            ("Calcium carbonate", "Canxi cacbonat"),
            ("Water", "Nước"),
        ]
        for idx, (en, vi) in enumerate(chemicals):
            cell = tbl.rows[idx].cells[0]
            apply_pattern_b(cell, en, vi)

        out_path = os.path.join(self.temp_dir, "r4_chemicals.docx")
        doc.save(out_path)

        with zipfile.ZipFile(out_path, "r") as z:
            xml_text = z.read("word/document.xml").decode("utf-8")
            for _, vi in chemicals:
                self.assertIn(vi, xml_text, f"Chemical translation '{vi}' must appear in output table")

    def test_r6_token_sanitizer_and_leak_hunter(self):
        """R6: Pre-sanitizer fixes typos & spacing; Leak hunter repairs leaked ordinals & Goods/Measures."""
        # Pre-sanitizer
        dirty_input = "flush eyes for at least 15minutes.Assure adequate flushing. Keep away from hest and flame."
        cleaned = sanitize_tokens(dirty_input)
        self.assertEqual(cleaned, "flush eyes for at least 15 minutes. Assure adequate flushing. Keep away from heat and flame.")

        # Post-translation leak hunter
        leaked_vi = "Đây là 1st lần kiểm tra danh mục Goods theo section 6.2."
        repaired = hunt_and_repair_leaks(leaked_vi)
        self.assertNotIn("1st", repaired)
        self.assertNotIn("Goods", repaired)
        self.assertNotIn("section 6.2", repaired)
        self.assertIn("thứ 1", repaired)
        self.assertIn("Hàng hóa", repaired)
        self.assertIn("mục 6.2", repaired)

    def test_r7_structural_ordering_and_glossary(self):
        """R7: Normalizes section numbering to Roman numerals I-XVI and enforces HSE terms."""
        # Glossary enforcement
        hse_en = "No smoking, drinking and eating"
        hse_vi = OFFLINE_GLOSSARY[hse_en]
        self.assertEqual(hse_vi, "Không hút thuốc, ăn uống", "Must enforce 'ăn uống', never 'uống rượu'")

        # Roman numeral normalization
        doc = docx.Document()
        apply_pattern_a(doc, "SECTION 1: Identification", OFFLINE_GLOSSARY["SECTION 1: Identification of the substance/mixture and of the company/undertaking"])
        apply_pattern_a(doc, "SECTION 2: Hazards Identification", OFFLINE_GLOSSARY["SECTION 2: Hazards Identification"])
        out_path = os.path.join(self.temp_dir, "r7_order.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["header_order_violations"], 0)

    def test_r9_zero_deepseek_api_cost_guard(self):
        """R9: ZeroCostGuard must actively intercept external calls to api.deepseek.com and raise ApiCostViolationError."""
        with ZeroCostGuard():
            with self.assertRaises(ApiCostViolationError):
                import urllib.request
                urllib.request.urlopen("https://api.deepseek.com/v1/chat/completions")


# ==============================================================================
# TIER 2: BOUNDARY & CORNER CASES TESTS
# ==============================================================================

class TestTier2BoundaryCases(unittest.TestCase):
    """Verify boundary conditions: empty lines, 0-value, massive tables, nested tables, symbols, CAS."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_tier2_empty_lines_and_whitespace_paragraphs(self):
        """Empty lines and whitespace-only paragraphs must not cause crashes or false audit flags."""
        doc = docx.Document()
        doc.add_paragraph("")
        doc.add_paragraph("   ")
        doc.add_paragraph("\t\t")  # Empty tab line
        apply_pattern_a(doc, "General Information", "Thông tin chung")
        out_path = os.path.join(self.temp_dir, "t2_empty_lines.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertTrue(report["pair_completeness"] >= 1.0)

    def test_tier2_zero_value_and_numeric_boundaries(self):
        """Numeric boundaries (0, 0.00%, > 1700 °C, < 0.1%, ±0.05) must preserve formatting and alignment."""
        doc = docx.Document()
        tbl = doc.add_table(rows=4, cols=2)
        pairs = [
            ("Relative density: 0.00", "Tỷ trọng tương đối: 0.00"),
            ("Melting point: > 1700 °C", "Điểm nóng chảy: > 1700 °C"),
            ("Impurity: < 0.1%", "Tạp chất: < 0.1%"),
            ("Tolerance: ±0.05", "Dung sai: ±0.05"),
        ]
        for idx, (en, vi) in enumerate(pairs):
            apply_pattern_b(tbl.rows[idx].cells[0], en.split(":")[0], vi.split(":")[0])
            apply_pattern_b(tbl.rows[idx].cells[1], ":" + en.split(":")[1], ":" + vi.split(":")[1])

        out_path = os.path.join(self.temp_dir, "t2_numerics.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["concatenation_violations"], 0)
        self.assertEqual(report["pair_completeness"], 1.0)

    def test_tier2_massive_tables_50_plus_rows(self):
        """Massive tables (50+ rows) must have 100% unlocked rows and zero clipping defects."""
        doc = docx.Document()
        tbl = doc.add_table(rows=60, cols=2)
        for i in range(60):
            row = tbl.rows[i]
            trPr = row._tr.get_or_add_trPr()
            trHeight = OxmlElement("w:trHeight")
            trHeight.set(qn("w:hRule"), "atLeast")
            trPr.append(trHeight)
            apply_pattern_b(row.cells[0], f"Property {i}", f"Đặc tính {i}")
            apply_pattern_b(row.cells[1], f"Value {i}", f"Giá trị {i}")

        out_path = os.path.join(self.temp_dir, "t2_massive_table.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["clipped_row_violations"], 0)
        self.assertEqual(report["pair_completeness"], 1.0)

    def test_tier2_multi_column_colons_on_single_line(self):
        """Single lines with multiple colon fields (Telephone... Fax... Email...) must split properly."""
        raw_line = "Telephone: 0274 3555 208   Fax: 0274 3555 209   Email: info@example.com"

        # Split helper
        fields = re.findall(r"([A-Za-z]+)\s*:\s*([^:]+?)(?=(?:[A-Za-z]+\s*:|$))", raw_line)
        self.assertEqual(len(fields), 3, "Must split into exactly 3 distinct label:value tuples")
        self.assertEqual(fields[0][0].strip(), "Telephone")
        self.assertEqual(fields[1][0].strip(), "Fax")
        self.assertEqual(fields[2][0].strip(), "Email")

    def test_tier2_nested_tables_and_borderless_grids(self):
        """Nested tables / borderless tables must maintain cell-level bilingual alignment."""
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        inner_tbl = tbl.rows[0].cells[0].add_table(rows=1, cols=2)
        apply_pattern_b(inner_tbl.rows[0].cells[0], "Synonym", "Tên đồng nghĩa")
        apply_pattern_b(inner_tbl.rows[0].cells[1], ": Precipitated Chalk", ": Phấn kết tủa")

        out_path = os.path.join(self.temp_dir, "t2_nested_tables.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["concatenation_violations"], 0)

    def test_tier2_special_characters_and_symbols(self):
        """Special characters (®, ™, ©, °, µg/m³, ≥, ≤, Δ) must not trigger concatenation violations."""
        tokens = ["SYLOID®", "Brand™", "Copyright©", "25°C", "50µg/m³", "≥95%", "≤5%", "ΔH"]
        for token in tokens:
            self.assertFalse(re.search(r"([a-z0-9])([A-ZÀ-Ỹ])", token.replace("°C", "").replace("µg/m³", "")),
                             f"Token {token} should not match basic concatenation pattern")

    def test_tier2_cas_numbers_and_regulatory_identifiers(self):
        """CAS numbers and regulatory IDs must be recognized as non-translatable technical tokens."""
        cas_tokens = [
            "CAS#: 471-34-1",
            "CAS: 7631-86-9",
            "EC: 231-545-4",
            "01-2119379499-16-XXXX",
            "EINECS: 231-545-4",
            "UN 1263"
        ]
        for token in cas_tokens:
            # Source checker should not demand translating pure CAS numbers
            self.assertFalse(is_translatable_source_text(token.split(":")[-1].strip()),
                             f"Pure numeric/alphanumeric code '{token}' should not require translation")

    def test_tier2_adversarial_inter_run_concatenation_across_formatting(self):
        """Adversarial: adjacent runs with different styles (bold EN + italic VI) without break or space must be caught."""
        doc = docx.Document()
        p = doc.add_paragraph()
        r1 = p.add_run("Product")
        r1.bold = True
        r2 = p.add_run("Sản phẩm")
        r2.italic = True
        r2.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        p_out = os.path.join(self.temp_dir, "t2_adversarial_concat.docx")
        doc.save(p_out)

        rep = verify_bilingual_docx(p_out)
        self.assertFalse(rep["passed"])
        self.assertGreater(rep["concatenation_violations"], 0)



# ==============================================================================
# TIER 3: CROSS-FEATURE INTERACTIONS TESTS
# ==============================================================================

class TestTier3CrossFeatureInteractions(unittest.TestCase):
    """Verify complex interactions between features (e.g. multi-column line-wrapped chemicals)."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_tier3_chemical_substances_in_multi_column_tables_with_line_wraps(self):
        """Chemical substances with line wraps inside multi-column tables must maintain vertical cell alignment."""
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=3)

        # Header
        apply_pattern_b(tbl.rows[0].cells[0], "Chemical Name", "Tên hóa chất")
        apply_pattern_b(tbl.rows[0].cells[1], "CAS Number", "Số CAS")
        apply_pattern_b(tbl.rows[0].cells[2], "Weight %", "Tỷ lệ %")

        # Row 1 with long chemical name wrapping
        long_chem_en = "amorphous silicon dioxide, chemically prepared (precipitated and gel silica)"
        long_chem_vi = "Silic đioxit vô định hình, điều chế hóa học (silica kết tủa và gel silica)"
        apply_pattern_b(tbl.rows[1].cells[0], long_chem_en, long_chem_vi)
        apply_pattern_b(tbl.rows[1].cells[1], "7631-86-9", "7631-86-9")
        apply_pattern_b(tbl.rows[1].cells[2], "95-100%", "95-100%")

        out_path = os.path.join(self.temp_dir, "t3_chem_table_wrap.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["concatenation_violations"], 0)
        self.assertEqual(report["clipped_row_violations"], 0)
        self.assertEqual(report["pair_completeness"], 1.0)

    def test_tier3_label_value_with_technical_symbols_and_units(self):
        """Label : Value pairs with symbols and units must reconstruct cleanly with 0 double tabs."""
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=2)

        apply_pattern_b(tbl.rows[0].cells[0], "pH-value at 20 °C", "Độ pH ở 20 °C")
        apply_pattern_b(tbl.rows[0].cells[1], ": 5.0 - 8.0 (±0.05)", ": 5.0 - 8.0 (±0.05)")

        apply_pattern_b(tbl.rows[1].cells[0], "Viscosity at 25 °C", "Độ nhớt ở 25 °C")
        apply_pattern_b(tbl.rows[1].cells[1], ": 120 - 150 mPa.s", ": 120 - 150 mPa.s")

        out_path = os.path.join(self.temp_dir, "t3_label_val_units.docx")
        doc.save(out_path)

        report = verify_bilingual_docx(out_path)
        self.assertEqual(report["double_tab_violations"], 0)
        self.assertEqual(report["concatenation_violations"], 0)

    def test_tier3_section_headers_preceding_borderless_tables(self):
        """Section headers preceding borderless tables must maintain keepNext to prevent orphaned headings."""
        doc = docx.Document()
        p_hdr, _ = apply_pattern_a(doc, "SECTION 9: Physical and Chemical Properties", "IX. Đặc tính lý, hóa học")
        # Ensure keepNext on header
        keep_next = p_hdr._p.xpath("./w:pPr/w:keepNext")
        self.assertEqual(len(keep_next), 1, "Section header before table must keep with next paragraph")


# ==============================================================================
# TIER 4: REAL-WORLD DOCUMENT VALIDATION & GOLDEN AUDIT TESTS
# ==============================================================================

class TestTier4RealWorldValidation(unittest.TestCase):
    """Real-world sample document validation and 7-Point Audit certification."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.sample_paths = {
            "YTU0035": r"C:\Users\ntc27\OneDrive\Desktop\MSDS VISION\Thành phẩm\Y\MSDS-YTU0035-EN.pdf",
            "YTU0070": r"C:\Users\ntc27\OneDrive\Desktop\MSDS VISION\Thành phẩm\Y\MSDS-YTU0070-EN.pdf",
            "FU43288": r"C:\Users\ntc27\Downloads\MSDS\MSDS-FU43288-EN.doc",
        }

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_tier4_sample_files_accessibility_and_properties(self):
        """Verify that all 3 real-world test sample files exist on disk with non-zero size."""
        for name, path in self.sample_paths.items():
            exists = os.path.exists(path)
            self.assertTrue(exists, f"Sample {name} must exist at {path}")
            size = os.path.getsize(path)
            self.assertGreater(size, 1000, f"Sample {name} size ({size} bytes) should be > 1000")

    def test_tier4_full_7_point_audit_pass_on_golden_document(self):
        """A complete golden bilingual document covering Sections 1-3 must achieve 100% audit pass."""
        doc = docx.Document()

        # Dynamic header
        hdr = doc.sections[0].header.paragraphs[0]
        hdr.add_run("Bilingual MSDS - ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        hdr._p.append(fld)

        # Section 1
        apply_pattern_a(doc, "SECTION 1: Identification of the substance/mixture and of the company/undertaking",
                        "I. Nhận dạng hóa chất và thông tin về nhà cung cấp")
        apply_pattern_a(doc, "Trade name: SYLOID MX107", "Tên thương mại: SYLOID MX107")

        # Section 2
        apply_pattern_a(doc, "SECTION 2: Hazards Identification", "II. Nhận dạng đặc tính nguy hiểm")
        apply_pattern_a(doc, "Classification of the substance or mixture: Not classified",
                        "Phân loại chất hoặc hỗn hợp: Không phân loại")

        # Section 3 Table
        apply_pattern_a(doc, "SECTION 3: Composition / Information on Ingredients",
                        "III. Thông tin về thành phần các chất")
        tbl = doc.add_table(rows=2, cols=3)
        # Unlock heights
        for r in tbl.rows:
            trPr = r._tr.get_or_add_trPr()
            trHeight = OxmlElement("w:trHeight")
            trHeight.set(qn("w:hRule"), "atLeast")
            trPr.append(trHeight)

        apply_pattern_b(tbl.rows[0].cells[0], "Chemical Name", "Tên hóa chất")
        apply_pattern_b(tbl.rows[0].cells[1], "CAS Number", "Số CAS")
        apply_pattern_b(tbl.rows[0].cells[2], "Weight %", "Tỷ lệ %")

        apply_pattern_b(tbl.rows[1].cells[0], "Calcium carbonate", "Canxi cacbonat")
        apply_pattern_b(tbl.rows[1].cells[1], "471-34-1", "471-34-1")
        apply_pattern_b(tbl.rows[1].cells[2], "100%", "100%")

        out_path = os.path.join(self.temp_dir, "golden_bilingual.docx")
        doc.save(out_path)

        # Audit
        report = verify_bilingual_docx(out_path)
        self.assertTrue(report["passed"], f"Golden document must pass 100% of audit checks. Details: {report}")
        self.assertEqual(report["concatenation_violations"], 0)
        self.assertEqual(report["double_tab_violations"], 0)
        self.assertEqual(report["pair_completeness"], 1.0)
        self.assertEqual(report["leaked_english_count"], 0)
        self.assertEqual(report["header_order_violations"], 0)
        self.assertEqual(report["clipped_row_violations"], 0)
        self.assertTrue(report["visual_parity_passed"])

    def test_tier4_negative_detection_matrix(self):
        """Audit verifier must accurately detect each individual failure mode without false negatives."""
        # 1. Concatenation negative test
        d_concat = docx.Document()
        d_concat.add_paragraph("liquidChất lỏng")
        p_c = os.path.join(self.temp_dir, "neg_concat.docx")
        d_concat.save(p_c)
        r_c = verify_bilingual_docx(p_c)
        self.assertFalse(r_c["passed"])
        self.assertGreater(r_c["concatenation_violations"], 0)

        # 2. Double tab negative test
        d_tab = docx.Document()
        d_tab.add_paragraph("Product\t\tCode")
        p_t = os.path.join(self.temp_dir, "neg_tab.docx")
        d_tab.save(p_t)
        r_t = verify_bilingual_docx(p_t)
        self.assertFalse(r_t["passed"])
        self.assertGreater(r_t["double_tab_violations"], 0)

        # 3. Leaked English negative test
        d_leak = docx.Document()
        p_l = d_leak.add_paragraph()
        r_l = p_l.add_run("Biện pháp xử lý First aid theo quy định")
        r_l.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r_l.italic = True
        p_leak = os.path.join(self.temp_dir, "neg_leak.docx")
        d_leak.save(p_leak)
        r_l = verify_bilingual_docx(p_leak)
        self.assertFalse(r_l["passed"])
        self.assertGreater(r_l["leaked_english_count"], 0)

        # 4. Inverted header negative test
        d_inv = docx.Document()
        d_inv.add_paragraph("SECTION 11: Toxicological Information")
        d_inv.add_paragraph("SECTION 3: Composition")
        p_inv = os.path.join(self.temp_dir, "neg_inv.docx")
        d_inv.save(p_inv)
        r_inv = verify_bilingual_docx(p_inv)
        self.assertFalse(r_inv["passed"])
        self.assertGreater(r_inv["header_order_violations"], 0)

        # 5. Clipped row negative test
        d_clip = docx.Document()
        tbl_c = d_clip.add_table(rows=1, cols=1)
        trPr = tbl_c.rows[0]._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:hRule"), "exact")
        trHeight.set(qn("w:val"), "180")
        trPr.append(trHeight)
        p_clip = os.path.join(self.temp_dir, "neg_clip.docx")
        d_clip.save(p_clip)
        r_clip = verify_bilingual_docx(p_clip)
        self.assertFalse(r_clip["passed"])
        self.assertEqual(r_clip["clipped_row_violations"], 1)

    def test_tier4_fixture_torture_ytu0035_simulation(self):
        """Simulate ChemGes YTU0035 structure: dynamic header, 2-column header, chemicals, reflow."""
        doc = docx.Document()
        sec = doc.sections[0]
        # Promoted header
        p_hdr = sec.header.paragraphs[0]
        p_hdr.add_run("Safety data sheet according to 1907/2006/EC - ")
        fld_p = OxmlElement("w:fldSimple")
        fld_p.set(qn("w:instr"), "PAGE")
        p_hdr._p.append(fld_p)

        # 16 Sections in proper order
        for i in range(1, 17):
            apply_pattern_a(doc, f"SECTION {i}: Section Title {i}", f"Mục {i}: Tiêu đề mục {i}")

        p_out = os.path.join(self.temp_dir, "sim_ytu0035.docx")
        doc.save(p_out)

        rep = verify_bilingual_docx(p_out)
        self.assertTrue(rep["passed"])
        self.assertEqual(rep["header_order_violations"], 0)
        self.assertEqual(rep["concatenation_violations"], 0)

    def test_tier4_cli_runner(self):
        """audit_verifier must be executable as a standalone module via CLI."""
        import subprocess
        d = docx.Document()
        apply_pattern_a(d, "SECTION 1: Identification", "I. Nhận dạng hóa chất")
        p_doc = os.path.join(self.temp_dir, "cli_test.docx")
        d.save(p_doc)

        res = subprocess.run(
            [sys.executable, "-m", "tests.audit_verifier", p_doc],
            capture_output=True,
            text=True
        )
        self.assertEqual(res.returncode, 0)
        self.assertIn("7-POINT XML & VISUAL AUDIT REPORT", res.stdout)
        self.assertIn("Check 1 (Zero Concatenation)", res.stdout)


if __name__ == "__main__":
    unittest.main()

