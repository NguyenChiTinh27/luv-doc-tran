"""
tests/test_context_profiler.py
──────────────────────────────
Bộ kiểm thử tự động cho Vòng 1: AI Context Profiler
- Kiểm tra trích xuất mẫu văn bản từ text, docx, và xlsx.
- Kiểm tra phân loại lĩnh vực chuyên ngành (Chemistry, Legal, Tech, Finance, General).
- Kiểm tra tạo mỏ neo bối cảnh (Context Anchor) và hồ sơ ngữ cảnh (ContextProfile).
- Kiểm tra cơ chế định hướng nghĩa từ vựng đa nghĩa theo ngành (solution -> dung dịch, compound -> hợp chất).
"""

import os
import sys
import unittest
import tempfile
import docx
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from core.context_profiler import AIContextProfiler, ContextProfile, DOMAIN_PRESETS


class TestContextProfiler(unittest.TestCase):

    def test_extract_sample_from_text(self):
        """Kiểm tra trích xuất mẫu từ văn bản thô theo số từ giới hạn."""
        sample_text = "Word " * 1000
        extracted = AIContextProfiler.extract_sample_from_text(sample_text, max_words=500)
        self.assertEqual(len(extracted.split()), 500)

    def test_extract_sample_from_docx(self):
        """Kiểm tra trích xuất mẫu từ tệp docx bao gồm tiêu đề, văn bản và bảng."""
        doc = docx.Document()
        doc.add_heading("Safety Data Sheet - Sulfuric Acid", level=1)
        doc.add_paragraph("Section 1: Identification of the substance and mixture. CAS Number: 7664-93-9.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Chemical Name"
        table.cell(0, 1).text = "Sulfuric Acid"
        table.cell(1, 0).text = "Concentration"
        table.cell(1, 1).text = "98% w/w"

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
            doc.save(tf.name)
            docx_path = tf.name

        try:
            sample = AIContextProfiler.extract_sample_from_docx(docx_path, max_words=100)
            self.assertIn("Safety Data Sheet", sample)
            self.assertIn("7664-93-9", sample)
            self.assertIn("Sulfuric Acid", sample)
        finally:
            if os.path.exists(docx_path):
                os.remove(docx_path)

    def test_extract_sample_from_xlsx(self):
        """Kiểm tra trích xuất mẫu từ tệp xlsx."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory"
        ws.append(["Item Code", "Chemical Description", "CAS Number", "Hazard Class"])
        ws.append(["CHEM-01", "Sodium Hydroxide Solution", "1310-73-2", "Corrosive 8"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            wb.save(tf.name)
            xlsx_path = tf.name

        try:
            sample = AIContextProfiler.extract_sample_from_xlsx(xlsx_path, max_words=100)
            self.assertIn("Chemical Description", sample)
            self.assertIn("1310-73-2", sample)
        finally:
            if os.path.exists(xlsx_path):
                os.remove(xlsx_path)

    def test_classify_domain_chemistry(self):
        """Kiểm tra nhận diện chuyên ngành Hóa chất / MSDS với số CAS và thuật ngữ an toàn."""
        text = "Material Safety Data Sheet for Sulfuric Acid solution. CAS: 7664-93-9. Hazard statements: causes severe skin burns."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "chemistry")
        self.assertGreaterEqual(conf, 0.7)

    def test_classify_domain_legal(self):
        """Kiểm tra nhận diện chuyên ngành Pháp lý / Hợp đồng thương mại."""
        text = "This Commercial Agreement is entered into by and between the Parties. In witness whereof, the signatories agree to all terms and governing law."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "legal")
        self.assertGreaterEqual(conf, 0.7)

    def test_classify_domain_tech(self):
        """Kiểm tra nhận diện chuyên ngành Công nghệ thông tin / Phần mềm."""
        text = "Microservices architecture deployment on Kubernetes server. Database configuration and API endpoint authentication protocol."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "tech")
        self.assertGreaterEqual(conf, 0.7)

    def test_classify_domain_finance(self):
        """Kiểm tra nhận diện chuyên ngành Tài chính / Kinh tế."""
        text = "Consolidated balance sheet and cash flow statement for fiscal year. Total assets, net revenue, EBITDA and liabilities."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "finance")
        self.assertGreaterEqual(conf, 0.7)

    def test_classify_domain_general(self):
        """Kiểm tra fallback về General khi văn bản không mang đặc trưng ngành rõ rệt."""
        text = "Welcome to the meeting. We are glad to see everyone today."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "general")

    def test_profile_document_creation_and_context_anchor(self):
        """Kiểm tra quy trình tạo ContextProfile và chuỗi Context Anchor hoàn chỉnh."""
        raw_text = "Section 2: Hazard Identification. Chemical compound and acidic solution. CAS Number: 7664-93-9."
        profile = AIContextProfiler.profile_document(raw_text)
        self.assertEqual(profile.domain_key, "chemistry")
        self.assertIn("Hóa chất", profile.domain)
        self.assertIn("solution", profile.glossary)
        self.assertEqual(profile.glossary["solution"], "dung dịch")
        self.assertIn("compound", profile.glossary)
        self.assertEqual(profile.glossary["compound"], "hợp chất")
        self.assertIn("BỐI CẢNH DỊCH THUẬT CHUYÊN NGÀNH", profile.context_anchor)
        self.assertIn("dung dịch", profile.context_anchor)

    def test_apply_glossary_disambiguation(self):
        """
        Kiểm tra xử lý từ ngữ đa nghĩa theo ngữ cảnh:
        - Trong ngành hóa chất: 'solution' -> 'dung dịch' (không phải 'giải pháp').
        - Trong ngành hóa chất: 'compound' -> 'hợp chất' (không phải 'khu phức hợp').
        - Trong ngành pháp lý: 'party' -> 'bên' (không phải 'bữa tiệc').
        """
        # 1. Hóa chất: solution dịch nhầm thành giải pháp -> tự sửa thành dung dịch
        chem_profile = ContextProfile(
            domain="Hóa chất",
            domain_key="chemistry",
            tone="Chuẩn mực",
            glossary=DOMAIN_PRESETS["chemistry"]["glossary"]
        )
        src1 = "Preparation of hydrochloric acid solution for laboratory analysis."
        trans1 = "Chuẩn bị giải pháp axit clohydric để phân tích trong phòng thí nghiệm."
        fixed1 = AIContextProfiler.apply_glossary_to_text(src1, trans1, chem_profile)
        self.assertIn("dung dịch axit clohydric", fixed1)
        self.assertNotIn("giải pháp", fixed1)

        # 2. Hóa chất: compound dịch nhầm thành khu phức hợp -> tự sửa thành hợp chất
        src2 = "This chemical compound exhibits strong stability."
        trans2 = "Khu phức hợp hóa học này thể hiện tính ổn định cao."
        fixed2 = AIContextProfiler.apply_glossary_to_text(src2, trans2, chem_profile)
        self.assertIn("Hợp chất hóa học", fixed2)
        self.assertNotIn("Khu phức hợp", fixed2)

        # 3. Pháp lý: party dịch nhầm thành bữa tiệc -> tự sửa thành bên
        legal_profile = ContextProfile(
            domain="Pháp lý",
            domain_key="legal",
            tone="Trang trọng",
            glossary=DOMAIN_PRESETS["legal"]["glossary"]
        )
        src3 = "Each party agrees to comply with the agreement."
        trans3 = "Mỗi bữa tiệc đồng ý tuân thủ hợp đồng."
        fixed3 = AIContextProfiler.apply_glossary_to_text(src3, trans3, legal_profile)
        self.assertIn("Mỗi bên đồng ý", fixed3)
        self.assertNotIn("bữa tiệc", fixed3)

    def test_classify_domain_german_chemical_msds(self):
        """Kiểm tra nhận diện chuyên ngành Hóa chất / MSDS từ tài liệu tiếng Đức có CAS-Nr."""
        text = "Sicherheitsdatenblatt gemäß Verordnung (EG) Nr. 1907/2006. Schwefelsäure 96%, CAS-Nr.: 7664-93-9. Gefahr: Verursacht schwere Verätzungen."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "chemistry")
        self.assertGreaterEqual(conf, 0.90)

    def test_classify_domain_french_chemical_msds(self):
        """Kiểm tra nhận diện chuyên ngành Hóa chất / MSDS từ tài liệu tiếng Pháp có No CAS."""
        text = "Fiche de données de sécurité selon le règlement (CE) no 1907/2006. Acide sulfurique 96%, No CAS 7664-93-9. Danger: Provoque de graves brûlures."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "chemistry")
        self.assertGreaterEqual(conf, 0.90)

    def test_classify_domain_isolated_cas_only(self):
        """Kiểm tra nhận diện chuyên ngành Hóa chất chỉ dựa vào sự hiện diện của số hiệu CAS."""
        text = "Sample item identifier with CAS 50-00-0 and UN 1830."
        domain, conf = AIContextProfiler.classify_domain_heuristic(text)
        self.assertEqual(domain, "chemistry")
        self.assertGreaterEqual(conf, 0.70)


if __name__ == "__main__":
    unittest.main()
