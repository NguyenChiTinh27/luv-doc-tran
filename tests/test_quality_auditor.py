"""
tests/test_quality_auditor.py
─────────────────────────────
Bộ kiểm thử tự động cho Vòng 3: AI Quality Audit & Auto-heal
- Kiểm tra trích xuất và bảo toàn mã số CAS và URL.
- Kiểm tra phát hiện và tự sửa lỗi (Auto-heal) đối với:
  1. Đoạn văn bị bỏ sót hoặc rỗng.
  2. Mã số CAS hoặc mã kỹ thuật bị rơi mất trong quá trình dịch.
  3. Placeholder __LUV_KEEP_XXXX__ chưa được khôi phục.
  4. Thuật ngữ sai lệch ngữ cảnh (đối chiếu hồ sơ ngữ cảnh Vòng 1).
- Kiểm tra tính hợp lệ của tệp Word (.docx) và Excel (.xlsx).
"""

import os
import sys
import unittest
import tempfile
import docx
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from core.context_profiler import ContextProfile, DOMAIN_PRESETS
from core.quality_auditor import AIQualityAuditor, AuditReport


class TestQualityAuditor(unittest.TestCase):

    def setUp(self):
        self.chem_profile = ContextProfile(
            domain="Hóa chất",
            domain_key="chemistry",
            tone="Khoa học",
            glossary=DOMAIN_PRESETS["chemistry"]["glossary"]
        )

    def test_extract_technical_tokens(self):
        """Kiểm tra trích xuất số hiệu CAS và URL kỹ thuật."""
        text = "Check Sulfuric Acid (CAS: 7664-93-9) and Sodium Hydroxide (CAS: 1310-73-2). Details at https://example.com/sds."
        tokens = AIQualityAuditor.extract_technical_tokens(text)
        self.assertIn("7664-93-9", tokens)
        self.assertIn("1310-73-2", tokens)
        self.assertIn("https://example.com/sds", tokens)

    def test_audit_clean_pairs_passes_with_perfect_score(self):
        """Kiểm tra cặp dịch thuật hoàn chỉnh đạt điểm toàn vẹn 100% không cần sửa lỗi."""
        pairs = [
            ("Sulfuric Acid (CAS: 7664-93-9)", "Axit sunfuric (CAS: 7664-93-9)"),
            ("Aqueous solution of chemical compound", "Dung dịch nước của hợp chất hóa học")
        ]
        healed, report = AIQualityAuditor.audit_and_heal_pairs(pairs, profile=self.chem_profile)
        self.assertEqual(len(healed), 2)
        self.assertTrue(report.audit_passed)
        self.assertEqual(report.integrity_score, 100.0)
        self.assertEqual(report.healed_segments, 0)
        self.assertEqual(report.missing_segments, 0)
        self.assertEqual(report.corrupted_tokens, 0)

    def test_audit_missing_translation_auto_heal(self):
        """Kiểm tra tự động sửa chữa khi phát hiện đoạn văn bản bị bỏ sót / dịch rỗng."""
        pairs = [
            ("Precautionary statement: Wear protective gloves.", "")
        ]
        healed, report = AIQualityAuditor.audit_and_heal_pairs(pairs, profile=self.chem_profile)
        self.assertEqual(len(healed), 1)
        self.assertGreater(len(healed[0].strip()), 0)
        self.assertEqual(report.healed_segments, 1)
        self.assertEqual(report.missing_segments, 1)

    def test_audit_corrupted_cas_token_auto_heal(self):
        """Kiểm tra phát hiện mất số hiệu CAS và tự động bù đắp token kỹ thuật."""
        pairs = [
            ("Chemical Identifier: CAS: 7664-93-9 for Sulfuric Acid", "Định danh hóa học: cho Axit Sunfuric")
        ]
        healed, report = AIQualityAuditor.audit_and_heal_pairs(pairs, profile=self.chem_profile)
        self.assertEqual(len(healed), 1)
        # Token 7664-93-9 phải được khôi phục vào bản dịch
        self.assertIn("7664-93-9", healed[0])
        self.assertEqual(report.healed_segments, 1)
        self.assertEqual(report.corrupted_tokens, 1)

    def test_audit_unrestored_placeholder_auto_heal(self):
        """Kiểm tra phát hiện và khôi phục placeholder __LUV_KEEP_0000__ còn sót lại."""
        pairs = [
            ("Contact: admin@example.com (CAS: 7664-93-9)", "Liên hệ: __LUV_KEEP_0000__ (CAS: 7664-93-9)")
        ]
        healed, report = AIQualityAuditor.audit_and_heal_pairs(pairs, profile=self.chem_profile)
        self.assertEqual(len(healed), 1)
        self.assertEqual(report.healed_segments, 1)

    def test_audit_terminology_mismatch_auto_heal(self):
        """Kiểm tra phát hiện và hiệu chỉnh thuật ngữ sai lệch ngữ cảnh (solution -> dung dịch)."""
        pairs = [
            ("Preparation of standard test solution", "Chuẩn bị giải pháp kiểm tra tiêu chuẩn")
        ]
        healed, report = AIQualityAuditor.audit_and_heal_pairs(pairs, profile=self.chem_profile)
        self.assertEqual(len(healed), 1)
        self.assertIn("dung dịch", healed[0])
        self.assertNotIn("giải pháp", healed[0])
        self.assertEqual(report.healed_segments, 1)

    def test_verify_docx_integrity_valid(self):
        """Kiểm tra xác thực tệp Word (.docx) hợp lệ."""
        doc = docx.Document()
        doc.add_paragraph("Kiểm tra tính toàn vẹn OpenXML.")
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
            doc.save(tf.name)
            docx_path = tf.name

        try:
            valid, err = AIQualityAuditor.verify_docx_integrity(docx_path)
            self.assertTrue(valid)
            self.assertIsNone(err)
        finally:
            if os.path.exists(docx_path):
                os.remove(docx_path)

    def test_verify_docx_integrity_corrupt(self):
        """Kiểm tra phát hiện tệp Word (.docx) bị hỏng cấu trúc ZIP / OpenXML."""
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
            tf.write(b"CORRUPTED_NOT_A_ZIP_HEADER")
            corrupt_path = tf.name

        try:
            valid, err = AIQualityAuditor.verify_docx_integrity(corrupt_path)
            self.assertFalse(valid)
            self.assertIsNotNone(err)
        finally:
            if os.path.exists(corrupt_path):
                os.remove(corrupt_path)

    def test_verify_xlsx_integrity_valid(self):
        """Kiểm tra xác thực tệp Excel (.xlsx) hợp lệ."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Test", 123])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            wb.save(tf.name)
            xlsx_path = tf.name

        try:
            valid, err = AIQualityAuditor.verify_xlsx_integrity(xlsx_path)
            self.assertTrue(valid)
            self.assertIsNone(err)
        finally:
            if os.path.exists(xlsx_path):
                os.remove(xlsx_path)

    def test_verify_xlsx_integrity_corrupt(self):
        """Kiểm tra phát hiện tệp Excel (.xlsx) bị hỏng cấu trúc ZIP / OpenXML."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            tf.write(b"FAKE_EXCEL_BYTES")
            corrupt_path = tf.name

        try:
            valid, err = AIQualityAuditor.verify_xlsx_integrity(corrupt_path)
            self.assertFalse(valid)
            self.assertIsNotNone(err)
        finally:
            if os.path.exists(corrupt_path):
                os.remove(corrupt_path)

    def test_extract_un_code_tokens(self):
        """Kiểm tra trích xuất mã UN vận chuyển hóa chất nguy hiểm (UN 1830, UN 1993)."""
        text = "Transport regulations: Sulfuric acid UN 1830 Class 8 and UN1993."
        tokens = AIQualityAuditor.extract_technical_tokens(text)
        self.assertTrue(any("1830" in t for t in tokens))
        self.assertTrue(any("1993" in t for t in tokens))

    def test_audit_corrupted_un_token_auto_heal(self):
        """Kiểm tra phát hiện và khôi phục mã UN bị thất lạc trong bản dịch."""
        pairs = [
            ("Sulfuric acid solution, UN 1830, Class 8", "Dung dịch axit sunfuric, Nhóm 8")
        ]
        healed, report = AIQualityAuditor.audit_and_heal_pairs(pairs, profile=self.chem_profile)
        self.assertEqual(len(healed), 1)
        self.assertIn("UN 1830", healed[0])
        self.assertEqual(report.healed_segments, 1)

    def test_unrestored_placeholder_stripped_cleanly(self):
        """Kiểm tra placeholder không khớp được dọn dẹp sạch sẽ không để lại rác trên văn bản."""
        pairs = [
            ("General note without codes", "Ghi chú chung __LUV_KEEP_9999__")
        ]
        healed, report = AIQualityAuditor.audit_and_heal_pairs(pairs, profile=self.chem_profile)
        self.assertEqual(len(healed), 1)
        self.assertNotIn("__LUV_KEEP_", healed[0])


if __name__ == "__main__":
    unittest.main()
