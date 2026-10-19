import os
import sys
import time
import subprocess
import unittest
from playwright.sync_api import sync_playwright

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARTIFACTS_DIR = r"C:\Users\ntc27\.gemini\antigravity\brain\9166a9f1-2c38-41c5-967b-79b51c7e5cf6"
PORT = 8092
SERVER_URL = f"http://127.0.0.1:{PORT}"

class TestE2EFrontend(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        # Clean previous state
        for fname in ["license.json", "trials_db.json", os.path.join("data", "license.sec"), os.path.join("data", "license.sec.tmp")]:
            fpath = os.path.join(BASE_DIR, fname)
            if os.path.exists(fpath):
                try:
                    os.remove(fpath)
                except Exception:
                    pass

        # Reset devices on test license in D1 if needed
        try:
            import urllib.request
            token = "rHHfK9eC_8lTH02SP0CQPnyXAqMMDUXxwV2MiPothM7Avc4TxNZgnkKp5CFsIBP_"
            req = urllib.request.Request(
                "https://luv-license-api.luv-sds.workers.dev/v1/admin/licenses/90a1d2c4-160f-4c87-b7f3-137a42eb1230/reset-devices",
                data=b"{}",
                headers={"Authorization": f"Bearer {token}", "User-Agent": "LUV-License-Admin/1.0", "Content-Type": "application/json"},
                method="POST"
            )
            urllib.request.urlopen(req, timeout=5)
        except Exception:
            pass

        # Start LUV-DOC-TRAN server on port 8092
        env = os.environ.copy()
        env["PORT"] = str(PORT)
        cls.server_proc = subprocess.Popen(
            [sys.executable, "-m", "uvicorn", "server:app", "--host", "127.0.0.1", "--port", str(PORT)],
            cwd=BASE_DIR,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            env=env
        )
        time.sleep(3)

    @classmethod
    def tearDownClass(cls):
        if cls.server_proc:
            try:
                if sys.platform == "win32":
                    subprocess.run(["taskkill", "/F", "/T", "/PID", str(cls.server_proc.pid)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                else:
                    cls.server_proc.terminate()
                cls.server_proc.wait(timeout=5)
            except Exception:
                pass

    def test_complete_frontend_flow(self):
        """Test full frontend UX: Dark Obsidian Studio branding, streamlined multi-file queue, per-file direct download, batch ZIP download, zero Split-View Studio, and zero PDF modals."""
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(viewport={"width": 1440, "height": 950})
            page = context.new_page()

            # 1. Visit homepage
            page.goto(SERVER_URL, timeout=15000)
            page.wait_for_selector(".brand-name")

            brand_text = page.inner_text(".brand-name").replace("\n", "")
            self.assertIn("LUV-DOC-TRAN", brand_text)

            # 2. Check Key Status Pill
            trial_text = page.inner_text("#trial-timer-text")
            self.assertTrue("Key" in trial_text or "Kích Hoạt" in trial_text or "kết nối" in trial_text)

            # 3. Check Tech Intro Bar
            tech_bar = page.inner_text(".tech-bar")
            self.assertIn("#002060", tech_bar)
            self.assertIn("Bảng biểu", tech_bar)

            # Take screenshot of Initial State
            ss_path1 = os.path.join(ARTIFACTS_DIR, "luv_doctran_e2e_home.png")
            page.screenshot(path=ss_path1)
            print(f"[E2E] Saved homepage screenshot: {ss_path1}")

            # 4. Open License Activation Modal (Single Key Input)
            page.click("#btn-open-license")
            page.wait_for_selector("#modal-license", state="visible")

            # Verify no registration name input exists
            self.assertIsNone(page.query_selector("#act-reg-name"), "#act-reg-name must be removed")

            # Fill in License Key (Genuine Cloudflare D1 Ed25519 Key)
            page.fill("#act-key", "LUV-YMZHK-QUKY8-D4ATT-TSJVY")

            # Click Activate
            page.click("#btn-submit-activate")
            page.wait_for_selector(".key-licensed", timeout=25000)
            page.wait_for_selector("#modal-license", state="hidden", timeout=10000)

            # Verify badge updated to valid
            header_badge = page.inner_text("#trial-pill")
            self.assertTrue("Key" in header_badge or "Hợp Lệ" in header_badge)

            # 5. Create 3 sample corporate files (Word + Excel) for Multi-file Upload testing (free of chemical keywords)
            sample_docx1 = os.path.join(BASE_DIR, "data", "temp", "e2e_sample_report.docx")
            sample_docx2 = os.path.join(BASE_DIR, "data", "temp", "e2e_sample_contract.docx")
            sample_xlsx = os.path.join(BASE_DIR, "data", "temp", "e2e_sample_sheet.xlsx")
            import docx
            for path, title, body in [
                (sample_docx1, "Báo Cáo Kế Hoạch Kinh Doanh & Tài Chính Doanh Nghiệp", "Tổng quan chiến lược phát triển thị trường và chỉ tiêu tăng trưởng năm 2026."),
                (sample_docx2, "Hợp Đồng Dịch Vụ Tư Vấn Công Nghệ Thông Tin", "Các điều khoản bảo mật, phạm vi công việc và lộ trình bàn giao giải pháp.")
            ]:
                d = docx.Document()
                d.add_heading(title, level=1)
                d.add_paragraph(body)
                d.save(path)

            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "KinhDoanh"
            ws["A1"] = "Khoản Mục"
            ws["B1"] = "Kế Hoạch Doanh Thu Năm 2026"
            ws["A2"] = "Dịch Vụ Số"
            ws["B2"] = "Tăng trưởng bền vững và hiệu quả cao"
            wb.save(sample_xlsx)

            # Multi-file Upload via file input (Word + Excel)
            page.set_input_files("#file-input", [sample_docx1, sample_docx2, sample_xlsx])
            time.sleep(0.5)

            # Check Queue Table is visible
            page.wait_for_selector("#queue-container", state="visible")
            badge_count = page.inner_text("#queue-count-badge")
            self.assertEqual(badge_count, "3 tệp")

            # Click Batch Processing Button
            page.click("#btn-start-batch")

            # Wait for all files in batch processing to finish
            page.wait_for_function("() => window.LuvApp && !window.LuvApp.isProcessingBatch", timeout=90000)
            page.wait_for_selector(".btn-row-word", state="visible", timeout=10000)
            page.wait_for_selector(".btn-row-excel", state="visible", timeout=10000)

            # Take screenshot of Queue with completed files
            ss_path_queue = os.path.join(ARTIFACTS_DIR, "luv_doctran_e2e_queue.png")
            page.screenshot(path=ss_path_queue)
            print(f"[E2E] Saved Queue screenshot: {ss_path_queue}")

            # 6. Verify Split-View Studio is completely eliminated from DOM (R1 & R2)
            self.assertIsNone(page.query_selector("#split-view-studio"), "#split-view-studio must not exist on DOM")
            self.assertIsNone(page.query_selector(".btn-row-view"), "btn-row-view (So Sánh) must not exist on DOM")
            self.assertIsNone(page.query_selector(".btn-row-compare"), "btn-row-compare must not exist on DOM")

            # 7. Verify per-file direct downloads (Word .docx and Excel .xlsx) & batch ZIP download (R2)
            word_btn = page.query_selector(".btn-row-word")
            self.assertIsNotNone(word_btn, "Per-file Word download button must exist in queue row")
            self.assertIn("Tải Word (.docx)", word_btn.inner_text())

            # Test clicking direct Word download button
            with page.expect_download(timeout=10000) as download_word_info:
                page.click(".btn-row-word")
            download_word = download_word_info.value
            self.assertTrue(download_word.suggested_filename.endswith(".docx"), "Downloaded file must be .docx")

            # Test clicking direct Excel download button
            excel_btn = page.query_selector(".btn-row-excel")
            self.assertIsNotNone(excel_btn, "Per-file Excel download button must exist in queue row")
            self.assertIn("Tải Excel (.xlsx)", excel_btn.inner_text())

            with page.expect_download(timeout=10000) as download_excel_info:
                page.click(".btn-row-excel")
            download_excel = download_excel_info.value
            self.assertTrue(download_excel.suggested_filename.endswith(".xlsx"), "Downloaded file must be .xlsx")

            # Check batch download all ZIP button is displayed and works
            btn_zip = page.query_selector("#btn-download-all")
            self.assertIsNotNone(btn_zip)
            self.assertIn("TẢI TOÀN BỘ TỆP ĐÃ DỊCH (.ZIP)", btn_zip.inner_text())

            with page.expect_download(timeout=15000) as download_zip_info:
                page.click("#btn-download-all")
            download_zip = download_zip_info.value
            self.assertTrue(download_zip.suggested_filename.endswith(".zip"), "Downloaded batch file must be .zip")

            # 8. Verify PDF Export is completely eliminated (R2)
            self.assertIsNone(page.query_selector("#modal-pdf-progress"), "PDF progress modal must be eliminated")
            self.assertIsNone(page.query_selector("#btn-studio-dl-pdf"), "PDF download button must be eliminated")

            # Take final full page screenshot of Streamlined Studio
            ss_path3 = os.path.join(ARTIFACTS_DIR, "luv_doctran_e2e_streamlined.png")
            page.screenshot(path=ss_path3)
            print(f"[E2E] Saved Streamlined Studio screenshot: {ss_path3}")

            browser.close()


if __name__ == "__main__":
    unittest.main()
