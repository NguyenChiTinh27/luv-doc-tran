"""
tests/test_engine_phase4_5.py
─────────────────────────────
Comprehensive unit tests for:
- Phase 4 Surgeon (unlock_table_rows, promote_headers_footers, normalize_tabs, apply_line_compaction)
- Phase 5 QC/QA (MasterQualityGate, GateResult, VisualQA, SelfHealingLoop)
"""

import os
import zipfile
import tempfile
import unittest
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from src.engine.phase4_surgeon import (
    unlock_table_rows,
    promote_headers_footers,
    normalize_tabs_and_indents,
    sanitize_package_xml,
    apply_line_compaction,
    OpenXmlSurgeon
)
from src.engine.phase5_qc_qa import (
    MasterQualityGate,
    GateResult,
    VisualQA,
    SelfHealingLoop
)
from tests.e2e_engine import apply_pattern_a, apply_pattern_b


class TestPhase4Surgeon(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_unlock_table_rows(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=2)
        # Lock row 0 with exact
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "300")
        trHeight.set(qn("w:hRule"), "exact")
        trPr.append(trHeight)

        unlocked = unlock_table_rows(doc)
        self.assertEqual(unlocked, 1)

        # Verify hRule changed to atLeast
        self.assertEqual(trHeight.get(qn("w:hRule")), "atLeast")
        # Verify cantSplit is present
        self.assertIsNotNone(trPr.find(qn("w:cantSplit")))

    def test_promote_headers_footers_and_purge_body(self):
        doc = docx.Document()
        doc.add_paragraph("First regular paragraph.")
        doc.add_paragraph("Page 1 of 5")  # Static pagination
        doc.add_paragraph("Another paragraph.")

        res = promote_headers_footers(doc, title_text="MSDS Document Title", remove_body_pagination=True)
        self.assertEqual(res["promoted_sections"], 1)
        self.assertEqual(res["removed_body_paragraphs"], 1)

        # Check body no longer contains "Page 1 of 5"
        body_texts = [p.text for p in doc.paragraphs]
        self.assertNotIn("Page 1 of 5", body_texts)

        # Save and verify header/footer in XML
        out_docx = os.path.join(self.temp_dir, "hf_test.docx")
        doc.save(out_docx)

        with zipfile.ZipFile(out_docx, "r") as z:
            header_xmls = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/header")]
            footer_xmls = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/footer")]

        self.assertTrue(any("MSDS Document Title" in h for h in header_xmls))
        self.assertTrue(any("PAGE" in f for f in footer_xmls))

    def test_normalize_tabs_and_indents(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        r = p.add_run("Label\t\tValue")
        cleaned = normalize_tabs_and_indents(doc)
        self.assertGreaterEqual(cleaned, 1)
        self.assertNotIn("\t\t", r.text)

    def test_apply_line_compaction(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        tbl.rows[0].cells[0].paragraphs[0].add_run("Compacted text content")

        compacted = apply_line_compaction(doc, compaction_factor=0.95)
        self.assertGreaterEqual(compacted, 1)

        # Verify w:w and w:spacing elements
        r_el = tbl.rows[0].cells[0].paragraphs[0].runs[0]._r
        rPr = r_el.find(qn("w:rPr"))
        self.assertIsNotNone(rPr)
        self.assertEqual(rPr.find(qn("w:w")).get(qn("w:val")), "94")
        self.assertEqual(rPr.find(qn("w:spacing")).get(qn("w:val")), "-4")


class TestPhase5QualityGate(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_master_quality_gate_on_golden_doc(self):
        doc = docx.Document()
        promote_headers_footers(doc, title_text="Golden Bilingual")
        apply_pattern_a(doc, "SECTION 1: Identification", "I. Nhận dạng sản phẩm và nhà sản xuất")
        apply_pattern_a(doc, "Product Name: Resin Super", "Tên sản phẩm: Nhựa Super")

        golden_path = os.path.join(self.temp_dir, "golden.docx")
        doc.save(golden_path)

        gate_res = MasterQualityGate.evaluate(golden_path)
        self.assertTrue(gate_res.passed)
        self.assertEqual(gate_res.concatenation_violations, 0)
        self.assertEqual(gate_res.double_tab_violations, 0)
        self.assertEqual(gate_res.pair_completeness, 1.0)
        self.assertIn("MASTER QUALITY GATE AUDIT RESULT", gate_res.summary_table())

    def test_self_healing_loop_execution(self):
        doc = docx.Document()
        promote_headers_footers(doc)
        apply_pattern_a(doc, "SECTION 1: Identification", "I. Nhận dạng")
        out_path = os.path.join(self.temp_dir, "healing.docx")
        doc.save(out_path)

        loop = SelfHealingLoop()
        res = loop.run_healing_loop(out_path, enable_visual_qa=False)
        self.assertTrue(res.passed)


if __name__ == "__main__":
    unittest.main()
