"""
tests/test_3stage_pipeline.py
─────────────────────────────
Bộ kiểm thử tích hợp toàn diện quy trình 3 giai đoạn AI (3-Stage AI Translation Pipeline):
- Vòng 1 (AI Context Profiler): Phân tích bối cảnh tài liệu đầu vào, lập hồ sơ ngữ cảnh (domain, tone, glossary).
- Vòng 2 (Context-Aware Translation): Động cơ dịch tích hợp mỏ neo bối cảnh, dịch chuẩn xác nghĩa chuyên ngành (solution -> dung dịch, compound -> hợp chất).
- Vòng 3 (AI Quality Audit & Auto-heal): Kiểm định chất lượng toàn diện, bảo toàn mã kỹ thuật / số CAS, tự sửa lỗi và xác nhận độ hợp lệ của tệp Word (.docx) & Excel (.xlsx).
"""

import os
import sys
import unittest
from unittest.mock import patch
import tempfile
from fastapi.testclient import TestClient
import docx
from docx.oxml.ns import qn
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from server import app
from core.license_client import LUVLicenseClient
from core.interlinear_translator import InterlinearTranslator
from core.excel_interlinear_translator import ExcelInterlinearTranslator
from core.context_profiler import AIContextProfiler
from core.quality_auditor import AIQualityAuditor


class Test3StagePipeline(unittest.TestCase):

    def setUp(self):
        self.client = TestClient(app)
        self.test_device = "TEST-DEV-3STAGE-001"
        self.patcher = patch.object(
            LUVLicenseClient,
            "check_access",
            return_value={"has_access": True, "access_type": "licensed", "tier": "Key Kích Hoạt (ENTERPRISE)", "revoked": False}
        )
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()

    def test_word_3stage_e2e_translation(self):
        """Kiểm tra quy trình 3 giai đoạn hoàn chỉnh cho tài liệu Word (.docx)."""
        doc = docx.Document()
        doc.add_heading("Safety Data Sheet - Sulfuric Acid", level=1)
        doc.add_paragraph("Chemical Identifier: Sulfuric Acid (CAS: 7664-93-9). Highly corrosive solution.")
        doc.add_paragraph("This chemical compound requires careful storage.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Property"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Flash Point"
        table.cell(1, 1).text = "Non-flammable"

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf_in:
            doc.save(tf_in.name)
            in_docx = tf_in.name

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf_out:
            out_docx = tf_out.name

        try:
            # Thực thi translate_docx
            res = InterlinearTranslator.translate_docx(
                docx_path=in_docx,
                output_path=out_docx,
                source_lang="en",
                target_lang="vi",
                output_mode="bilingual",
                provider="google"
            )

            # 1. Kiểm tra Vòng 1: AI Context Profiler
            self.assertIn("context_profile", res)
            cp = res["context_profile"]
            self.assertEqual(cp["domain_key"], "chemistry")
            self.assertIn("Hóa chất", cp["domain"])
            self.assertIn("solution", cp["glossary"])
            self.assertEqual(cp["glossary"]["solution"], "dung dịch")

            # 2. Kiểm tra Vòng 2: Context-Aware Translation
            # Tệp đầu ra phải chứa bản dịch với "dung dịch" và "hợp chất"
            out_doc = docx.Document(out_docx)
            full_text = " ".join([p.text for p in out_doc.paragraphs])
            self.assertIn("dung dịch", full_text.lower())
            self.assertIn("hợp chất", full_text.lower())

            # Kiểm tra màu sắc #002060 và italic cho dòng dịch
            has_blue_italic = False
            for p in out_doc.paragraphs:
                for r in p.runs:
                    if r.italic and r.font.color and r.font.color.rgb:
                        color_hex = str(r.font.color.rgb)
                        if color_hex.upper() == "002060":
                            has_blue_italic = True
                            break
            self.assertTrue(has_blue_italic, "Dòng dịch phải mang định dạng màu xanh đậm #002060 và in nghiêng.")

            # 3. Kiểm tra Vòng 3: AI Quality Audit & Auto-heal
            self.assertIn("audit_report", res)
            ar = res["audit_report"]
            self.assertTrue(ar["audit_passed"])
            self.assertTrue(ar["file_valid"])
            self.assertGreaterEqual(ar["integrity_score"], 90.0)

            # Mã CAS 7664-93-9 phải được bảo toàn 100% trong văn bản kết quả
            self.assertIn("7664-93-9", full_text)

            # Metadata 3 giai đoạn hoàn tất
            self.assertEqual(len(res["stages_completed"]), 3)

        finally:
            if os.path.exists(in_docx):
                os.remove(in_docx)
            if os.path.exists(out_docx):
                os.remove(out_docx)

    def test_excel_3stage_e2e_translation(self):
        """Kiểm tra quy trình 3 giai đoạn hoàn chỉnh cho bảng tính Excel (.xlsx)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Chemical Inventory"
        ws.append(["Item ID", "Chemical Description", "CAS Number", "Quantity", "Formula Total"])
        ws.append(["CHEM-001", "Concentrated Hydrochloric Acid Solution", "7647-01-0", 50, "=SUM(D2:D2)"])
        ws.append(["CHEM-002", "Sodium Hydroxide Compound", "1310-73-2", 30, "=D3*2"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf_in:
            wb.save(tf_in.name)
            in_xlsx = tf_in.name

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf_out:
            out_xlsx = tf_out.name

        try:
            res = ExcelInterlinearTranslator.translate_excel(
                input_path=in_xlsx,
                output_path=out_xlsx,
                source_lang="en",
                target_lang="vi",
                provider="google"
            )

            # 1. Kiểm tra Vòng 1
            self.assertIn("context_profile", res)
            cp = res["context_profile"]
            self.assertEqual(cp["domain_key"], "chemistry")

            # 2. Kiểm tra Vòng 2: CellRichText
            out_wb = openpyxl.load_workbook(out_xlsx, data_only=False)
            out_ws = out_wb.active

            # Kiểm tra bảo toàn công thức toán học (=SUM...)
            self.assertEqual(str(out_ws["E2"].value).strip(), "=SUM(D2:D2)")
            self.assertEqual(str(out_ws["E3"].value).strip(), "=D3*2")

            # Kiểm tra ô mô tả hóa chất có dòng song ngữ
            val_b2 = out_ws["B2"].value
            self.assertIsNotNone(val_b2)
            cell_str = str(val_b2)
            self.assertIn("Hydrochloric Acid", cell_str)
            self.assertIn("dung dịch", cell_str.lower())

            # 3. Kiểm tra Vòng 3
            self.assertIn("audit_report", res)
            ar = res["audit_report"]
            self.assertTrue(ar["audit_passed"])
            self.assertTrue(ar["file_valid"])
            self.assertIn("7647-01-0", str(out_ws["C2"].value))

            out_wb.close()

        finally:
            if os.path.exists(in_xlsx):
                os.remove(in_xlsx)
            if os.path.exists(out_xlsx):
                os.remove(out_xlsx)

    def test_api_translate_endpoint_returns_3stage_metadata(self):
        """Kiểm tra API /api/translate trả về đầy đủ metadata 3 giai đoạn AI cho giao diện người dùng."""
        doc = docx.Document()
        doc.add_heading("Commercial Contract & Agreement", level=1)
        doc.add_paragraph("This Agreement is entered into by and between the Parties. Governing law shall apply.")

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
            doc.save(tf.name)
            docx_path = tf.name

        try:
            with open(docx_path, "rb") as f:
                resp = self.client.post(
                    "/api/translate",
                    files={"file": ("contract_sample.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                    data={
                        "source_lang": "en",
                        "target_lang": "vi",
                        "output_mode": "bilingual",
                        "provider": "google",
                        "engine_mode": "auto",
                        "device_id": self.test_device
                    }
                )

            self.assertEqual(resp.status_code, 200)
            res = resp.json()
            self.assertTrue(res.get("success"))

            # Kiểm tra 3 giai đoạn
            self.assertIn("context_profile", res)
            self.assertEqual(res["context_profile"]["domain_key"], "legal")

            self.assertIn("audit_report", res)
            self.assertTrue(res["audit_report"]["audit_passed"])
            self.assertTrue(res["audit_report"]["file_valid"])

            self.assertIn("stages_completed", res)
            self.assertEqual(len(res["stages_completed"]), 3)

            # Kiểm tra tệp Word (.docx) ưu tiên tải về
            self.assertIn("download_docx_url", res)
            self.assertEqual(res.get("primary_format"), "docx")

        finally:
            if os.path.exists(docx_path):
                os.remove(docx_path)

    def test_excel_cas_and_un_token_integrity_without_duplication(self):
        """Kiểm tra bảng tính Excel chứa mã CAS và UN dịch hoàn chỉnh, đạt 100% điểm toàn vẹn và không bị lặp token."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "MSDS Inventory"
        ws.append(["Product Name", "Technical Spec"])
        ws.append(["Sulfuric acid 98% CAS 7664-93-9", "Dangerous Goods UN 1830 Solution"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf_in:
            wb.save(tf_in.name)
            in_xlsx = tf_in.name

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf_out:
            out_xlsx = tf_out.name

        try:
            res = ExcelInterlinearTranslator.translate_excel(
                input_path=in_xlsx,
                output_path=out_xlsx,
                source_lang="en",
                target_lang="vi",
                provider="google"
            )
            self.assertTrue(res["success"])
            ar = res["audit_report"]
            self.assertTrue(ar["audit_passed"])
            self.assertTrue(ar["file_valid"])
            self.assertEqual(ar["corrupted_tokens"], 0)
            self.assertEqual(ar["integrity_score"], 100.0)

            # Đọc lại workbook và kiểm tra nội dung ô
            out_wb = openpyxl.load_workbook(out_xlsx, data_only=False)
            out_ws = out_wb.active
            cell_val = str(out_ws["A2"].value)
            self.assertIn("7664-93-9", cell_val)
            # Mã CAS không được lặp lại dạng "(7664-93-9)"
            self.assertNotIn("(7664-93-9)", cell_val)
            self.assertNotIn("__LUV_KEEP_", cell_val)

            cell_b2 = str(out_ws["B2"].value)
            self.assertIn("UN 1830", cell_b2)
            self.assertNotIn("__LUV_KEEP_", cell_b2)
            out_wb.close()
        finally:
            if os.path.exists(in_xlsx):
                os.remove(in_xlsx)
            if os.path.exists(out_xlsx):
                os.remove(out_xlsx)


if __name__ == "__main__":
    unittest.main()
