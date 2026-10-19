"""
Independent Forensic Verification Script by Forensic Auditor.
Empirically tests and stress-tests all R1-R9 claims in core/bilingual_qc_v2.py:
1. Pattern A & Pattern B OpenXML structure
2. Dynamic row height unlocker (exact -> atLeast)
3. 2-column borderless grid with 28%/72% column widths
4. Token sanitizer, leak hunter, structure normalizer
5. Zero API cost guard & offline rule-based synthesis
6. Negative detection matrix (verifying audit verifier catches all defects)
7. End-to-end document reconstruction and 7-point audit
"""

import os
import sys
import tempfile
import zipfile
import lxml.etree as etree

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.bilingual_qc_v2 import (
    BilingualQCReconstructionEngine,
    reconstruct_bilingual_document,
    verify_bilingual_docx,
    ZeroCostGuard,
    ApiCostViolationError,
    TokenSanitizer,
    LeakHunter,
    RuleBasedSynthesizer,
    PairingBuilder,
    ReflowUnlocker,
    GridReconstructor,
    ParagraphConsolidator,
    StructureNormalizer,
    W_NS,
    NSMAP,
)


def run_all_forensic_audits():
    print("=== STARTING INDEPENDENT FORENSIC AUDIT ===")

    # -------------------------------------------------------------
    # 1. Zero API Cost Guard Verification
    # -------------------------------------------------------------
    print("\n[Check 1] Zero API Cost Guard Verification...")
    with ZeroCostGuard():
        # Verify that paid endpoints raise ApiCostViolationError
        import urllib.request
        try:
            urllib.request.urlopen("https://api.deepseek.com/chat/completions")
            assert False, "FAILED: ZeroCostGuard did not block api.deepseek.com urllib request!"
        except ApiCostViolationError as e:
            print(f"  OK: Successfully blocked urllib call to deepseek: {e}")

        try:
            import requests
            requests.get("https://api.deepseek.com/v1/models")
            assert False, "FAILED: ZeroCostGuard did not block api.deepseek.com requests call!"
        except ApiCostViolationError as e:
            print(f"  OK: Successfully blocked requests call to deepseek: {e}")

    # -------------------------------------------------------------
    # 2. Pattern A Dual-Paragraph OpenXML Construction
    # -------------------------------------------------------------
    print("\n[Check 2] Pattern A Dual-Paragraph Construction...")
    doc = docx.Document()
    p_src, p_trans = PairingBuilder.apply_pattern_a(
        parent_container=doc,
        en_text="First aid measures",
        vi_text="Biện pháp sơ cứu về y tế",
        font_name="Times New Roman",
        font_size_pt=10.5,
        is_bold=True
    )
    # Verify p_src
    pPr_src = p_src._p.find(qn("w:pPr"))
    assert pPr_src is not None, "Missing pPr on source paragraph"
    assert pPr_src.find(qn("w:keepNext")) is not None, "Missing <w:keepNext/> on source paragraph"
    sp_src = pPr_src.find(qn("w:spacing"))
    assert sp_src is not None and sp_src.get(qn("w:after")) == "0", "Source space_after is not 0"

    # Verify p_trans
    pPr_trans = p_trans.find(qn("w:pPr"))
    assert pPr_trans is not None, "Missing pPr on trans paragraph"
    sp_trans = pPr_trans.find(qn("w:spacing"))
    assert sp_trans is not None and sp_trans.get(qn("w:after")) == "80", "Trans space_after is not 80 (4pt)"
    
    # Check color and italic
    color_el = p_trans.xpath(".//w:color")
    assert color_el and color_el[0].get(qn("w:val")) == "1B365D", "Trans color is not 1B365D"
    italic_el = p_trans.xpath(".//w:i")
    assert bool(italic_el), "Trans paragraph is not italic"
    print("  OK: Pattern A generates valid OpenXML <w:keepNext/>, space_after=0, space_after=80, #1B365D, <w:i/>.")

    # -------------------------------------------------------------
    # 3. Pattern B Single-Cell OpenXML Construction
    # -------------------------------------------------------------
    print("\n[Check 3] Pattern B Single-Cell Construction...")
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    PairingBuilder.apply_pattern_b(
        cell=cell,
        en_text="Silicon dioxide",
        vi_text="Silic đioxit"
    )
    # Verify vertical centering
    tcPr = cell._tc.find(qn("w:tcPr"))
    assert tcPr is not None, "Missing tcPr"
    vAlign = tcPr.find(qn("w:vAlign"))
    assert vAlign is not None and vAlign.get(qn("w:val")) == "center", "Missing vertical center alignment"

    # Verify explicit <w:br/> between runs
    p = cell.paragraphs[0]
    runs = p.runs
    assert len(runs) >= 3, f"Expected at least 3 runs, got {len(runs)}"
    has_br = bool(cell._tc.xpath(".//w:br"))
    assert has_br, "Missing explicit <w:br/> in middle run"
    assert runs[-1].font.italic is True, "Vietnamese run must be italic"
    assert runs[-1].font.color.rgb == docx.shared.RGBColor(0x1B, 0x36, 0x5D), "Vietnamese run color must be #1B365D"
    print("  OK: Pattern B generates vertical center, explicit <w:r><w:br/></w:r>, and styled #1B365D italic run.")

    # -------------------------------------------------------------
    # 4. Dynamic Reflow Unlocker (exact -> atLeast)
    # -------------------------------------------------------------
    print("\n[Check 4] Dynamic Reflow Unlocker...")
    doc_reflow = docx.Document()
    t = doc_reflow.add_table(rows=2, cols=2)
    # Add exact height to rows
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "360")
        trHeight.set(qn("w:hRule"), "exact")
        trPr.append(trHeight)

    # Count exact before
    exact_before = len(doc_reflow.element.xpath("//w:trHeight[@w:hRule='exact']"))
    assert exact_before == 2, f"Expected 2 exact rows, got {exact_before}"

    ReflowUnlocker.unlock_table_row_heights(doc_reflow)

    exact_after = len(doc_reflow.element.xpath("//w:trHeight[@w:hRule='exact']"))
    atleast_after = len(doc_reflow.element.xpath("//w:trHeight[@w:hRule='atLeast']"))
    cantsplit_after = len(doc_reflow.element.xpath("//w:cantSplit"))
    assert exact_after == 0, f"Expected 0 exact rows, got {exact_after}"
    assert atleast_after == 2, f"Expected 2 atLeast rows, got {atleast_after}"
    assert cantsplit_after >= 2, f"Expected cantSplit on rows, got {cantsplit_after}"
    print("  OK: ReflowUnlocker successfully converted all exact row heights to atLeast and added cantSplit.")

    # -------------------------------------------------------------
    # 5. Invisible 2-Column Grid Reconstruction (28% / 72%)
    # -------------------------------------------------------------
    print("\n[Check 5] 2-Column Grid Reconstruction...")
    doc_grid = docx.Document()
    doc_grid.add_paragraph("Product name \t : Silicon Dioxide")
    doc_grid.add_paragraph("CAS Number \t : 7631-86-9")
    doc_grid.add_paragraph("Telephone: 012345678 Fax: 098765432")

    GridReconstructor.reconstruct_label_value_blocks(doc_grid)

    tables = doc_grid.tables
    assert len(tables) >= 1, "Expected at least 1 reconstructed table"
    reconstructed_tbl = tables[0]._tbl
    # Check borders none
    borders = reconstructed_tbl.xpath(".//w:tblBorders/*")
    assert borders, "Missing tblBorders"
    for b in borders:
        assert b.get(qn("w:val")) == "none", f"Border {b.tag} is not none"

    # Check grid col widths: 2520 (28%) and 6480 (72%)
    grid_cols = reconstructed_tbl.xpath(".//w:tblGrid/w:gridCol")
    assert len(grid_cols) == 2, f"Expected 2 gridCols, got {len(grid_cols)}"
    assert grid_cols[0].get(qn("w:w")) == "2520", f"Col 0 width is {grid_cols[0].get(qn('w:w'))}, expected 2520"
    assert grid_cols[1].get(qn("w:w")) == "6480", f"Col 1 width is {grid_cols[1].get(qn('w:w'))}, expected 6480"
    print("  OK: GridReconstructor built borderless table with exact 2520 (28%) and 6480 (72%) column widths.")

    # -------------------------------------------------------------
    # 6. Token Sanitizer, Leak Hunter & Structural Ordering
    # -------------------------------------------------------------
    print("\n[Check 6] Token Sanitizer, Leak Hunter & Structural Ordering...")
    # Sanitizer
    s1 = TokenSanitizer.sanitize("Wash for 15minutes.Assure adequate flushing.")
    assert "15 minutes. Assure" in s1, f"Failed unit/period spacing: {s1}"
    s2 = TokenSanitizer.sanitize("keep away from hest and flame and form ignition sources")
    assert "heat and flame" in s2 and "from ignition sources" in s2, f"Failed OCR typo fix: {s2}"
    print("  OK: TokenSanitizer regex fixes unit spacing, period spacing, and OCR typos.")

    # Leak Hunter
    lh1 = LeakHunter.hunt_and_repair("Biện pháp ở 1st cấp độ và kiểm soát Goods theo section6.2")
    assert "thứ 1" in lh1, f"Failed ordinal repair: {lh1}"
    assert "Hàng hóa" in lh1, f"Failed Goods repair: {lh1}"
    assert "mục 6.2" in lh1, f"Failed section ref repair: {lh1}"
    print("  OK: LeakHunter successfully repaired ordinals, keywords, and section references.")

    # Structural Ordering
    doc_struct = docx.Document()
    doc_struct.add_paragraph("SECTION 4: First Aid Measures")
    doc_struct.add_paragraph("Details of section 4...")
    doc_struct.add_paragraph("SECTION 2: Hazards Identification")
    doc_struct.add_paragraph("Details of section 2...")
    doc_struct.add_paragraph("SECTION 1: Identification")
    doc_struct.add_paragraph("Details of section 1...")

    StructureNormalizer.repair_inverted_headers(doc_struct)
    StructureNormalizer.normalize_section_headers(doc_struct)

    p_texts = [p.text for p in doc_struct.paragraphs if p.text]
    # Check section order
    sec_headers = [t for t in p_texts if t.startswith("I.") or t.startswith("II.") or t.startswith("IV.")]
    assert sec_headers == [
        "I. Identification of the substance/mixture and of the company/undertaking",
        "II. Hazards Identification",
        "IV. First Aid Measures"
    ], f"Failed section reordering: {sec_headers}"
    print("  OK: StructureNormalizer reordered inverted sections into strictly ascending Roman order.")

    # -------------------------------------------------------------
    # 7. Negative Detection Matrix (Verifying Audit Verifier Authenticity)
    # -------------------------------------------------------------
    print("\n[Check 7] Negative Detection Matrix on Audit Verifier...")
    with tempfile.TemporaryDirectory() as tmpdir:
        # Defect 1: Concatenation
        doc_bad1 = docx.Document()
        p_bad1 = doc_bad1.add_paragraph()
        r1 = p_bad1.add_run("Identification")
        r2 = p_bad1.add_run("Section")  # concatenation without break or space
        p1_path = os.path.join(tmpdir, "bad1.docx")
        doc_bad1.save(p1_path)
        res1 = verify_bilingual_docx(p1_path)
        assert res1["concatenation_violations"] > 0, "Audit failed to catch concatenation violation!"
        assert res1["passed"] is False, "Audit falsely passed bad concatenation!"
        print("  OK: Negative test caught concatenation violation.")

        # Defect 2: Double Tab
        doc_bad2 = docx.Document()
        doc_bad2.add_paragraph("Label \t\t Value")
        p2_path = os.path.join(tmpdir, "bad2.docx")
        doc_bad2.save(p2_path)
        res2 = verify_bilingual_docx(p2_path)
        assert res2["double_tab_violations"] > 0, "Audit failed to catch double tab!"
        assert res2["passed"] is False, "Audit falsely passed double tab!"
        print("  OK: Negative test caught double-tab violation.")

        # Defect 3: Exact table row height
        doc_bad3 = docx.Document()
        t_bad = doc_bad3.add_table(rows=1, cols=1)
        trPr = t_bad.rows[0]._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "360")
        trHeight.set(qn("w:hRule"), "exact")
        trPr.append(trHeight)
        p3_path = os.path.join(tmpdir, "bad3.docx")
        doc_bad3.save(p3_path)
        res3 = verify_bilingual_docx(p3_path)
        assert res3["clipped_row_violations"] > 0, "Audit failed to catch exact row height!"
        assert res3["passed"] is False, "Audit falsely passed exact row height!"
        print("  OK: Negative test caught exact row height violation.")

        # Defect 4: Missing keepNext on source body line
        doc_bad4 = docx.Document()
        p_src_bad = doc_bad4.add_paragraph("First aid measures")
        # Notice: no keepNext
        p_vi_bad = doc_bad4.add_paragraph("Biện pháp sơ cứu")
        p_vi_bad.runs[0].font.italic = True
        p_vi_bad.runs[0].font.color.rgb = docx.shared.RGBColor(0x1B, 0x36, 0x5D)
        p4_path = os.path.join(tmpdir, "bad4.docx")
        doc_bad4.save(p4_path)
        res4 = verify_bilingual_docx(p4_path)
        assert res4["visual_parity_passed"] is False, "Audit failed to catch missing keepNext!"
        print("  OK: Negative test caught missing keepNext on source line.")

    # -------------------------------------------------------------
    # 8. Full End-to-End Reconstruction & 7-Point Audit
    # -------------------------------------------------------------
    print("\n[Check 8] Full End-to-End Reconstruction & 7-Point Audit...")
    with tempfile.TemporaryDirectory() as tmpdir:
        input_doc = os.path.join(tmpdir, "input.docx")
        output_doc = os.path.join(tmpdir, "output_bilingual.docx")

        d_in = docx.Document()
        d_in.add_paragraph("MATERIAL SAFETY DATA SHEET")
        d_in.add_paragraph("SECTION 1: Identification of the substance/mixture and of the company/undertaking")
        d_in.add_paragraph("Product name : Silicon dioxide")
        d_in.add_paragraph("CAS Number : 7631-86-9")
        d_in.add_paragraph("SECTION 2: Hazards Identification")
        d_in.add_paragraph("No smoking, drinking and eating. Avoid contact with skin.")
        tbl_chem = d_in.add_table(rows=2, cols=2)
        tbl_chem.rows[0].cells[0].text = "Substance Name"
        tbl_chem.rows[0].cells[1].text = "Concentration (%)"
        tbl_chem.rows[1].cells[0].text = "Silicon dioxide"
        tbl_chem.rows[1].cells[1].text = "95-100%"
        # Add exact height to verify reflow unlocker
        trPr = tbl_chem.rows[0]._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "360")
        trHeight.set(qn("w:hRule"), "exact")
        trPr.append(trHeight)

        d_in.save(input_doc)

        out_path = reconstruct_bilingual_document(
            input_path=input_doc,
            output_path=output_doc,
            zero_cost=True
        )
        assert os.path.exists(out_path), "Output document was not created"

        report = verify_bilingual_docx(out_path)
        print(f"  Audit Report: Passed={report['passed']}")
        print(f"  - Concat: {report['concatenation_violations']}")
        print(f"  - Double Tab: {report['double_tab_violations']}")
        print(f"  - Completeness: {report['pair_completeness'] * 100:.1f}%")
        print(f"  - Leaked English: {report['leaked_english_count']}")
        print(f"  - Header Order: {report['header_order_violations']}")
        print(f"  - Clipped Rows: {report['clipped_row_violations']}")
        print(f"  - Visual Parity: {report['visual_parity_passed']}")

        assert report["passed"] is True, f"7-Point audit failed on reconstructed document: {report['details']}"
        assert report["concatenation_violations"] == 0
        assert report["double_tab_violations"] == 0
        assert report["pair_completeness"] == 1.0
        assert report["leaked_english_count"] == 0
        assert report["header_order_violations"] == 0
        assert report["clipped_row_violations"] == 0
        assert report["visual_parity_passed"] is True
        print("  OK: End-to-end reconstruction passed 100% of 7-Point XML & Visual Audit!")

    print("\n=== ALL FORENSIC AUDIT CHECKS PASSED EMPIRICALLY ===")


if __name__ == "__main__":
    run_all_forensic_audits()
