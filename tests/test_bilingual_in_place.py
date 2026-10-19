"""
tests/test_bilingual_in_place.py
────────────────────────────────
Bộ kiểm thử tự động xác thực tính năng:
Tự động nhận diện tài liệu đã là song ngữ và thay thế bản dịch tại chỗ
vào đúng phần ngôn ngữ thứ 2 của file đó (In-Place Secondary Language Replacement).
"""

import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import docx
from docx import Document
from docx.shared import Pt, RGBColor
import openpyxl
from openpyxl.cell.rich_text import CellRichText

from core.interlinear_translator import (
    detect_script_type,
    is_different_language_pair,
    is_paragraph_styled_as_translation,
    InterlinearTranslator,
    DARK_BLUE,
)
from core.excel_interlinear_translator import ExcelInterlinearTranslator


class TestScriptAndLanguageDetection(unittest.TestCase):
    """Kiểm thử thuật toán nhận diện hệ chữ và cặp ngôn ngữ song ngữ."""

    def test_detect_script_type(self):
        # Latin ASCII
        self.assertEqual(detect_script_type("Product Name"), "latin_ascii")
        self.assertEqual(detect_script_type("Section 1: Identification"), "latin_ascii")
        self.assertEqual(detect_script_type("CAS 67-64-1"), "latin_ascii")

        # CJK (Trung, Nhật, Hàn)
        self.assertEqual(detect_script_type("产品名称"), "cjk")
        self.assertEqual(detect_script_type("第一部分：化学品标识"), "cjk")
        self.assertEqual(detect_script_type("製品名"), "cjk")
        self.assertEqual(detect_script_type("제품명"), "cjk")

        # Vietnamese (Quốc ngữ có dấu)
        self.assertEqual(detect_script_type("Tên sản phẩm"), "vietnamese")
        self.assertEqual(detect_script_type("Nhận dạng hóa chất"), "vietnamese")
        self.assertEqual(detect_script_type("Cảnh báo nguy cơ"), "vietnamese")

        # Cyrillic (Nga)
        self.assertEqual(detect_script_type("Идентификация"), "cyrillic")

        # Arabic & Thai
        self.assertEqual(detect_script_type("معلومات السلامة"), "arabic")
        self.assertEqual(detect_script_type("ข้อมูลความปลอดภัย"), "thai")

        # Accented European Latin
        self.assertEqual(detect_script_type("Identificación del producto"), "latin_accented")

        # Other (số, ký hiệu)
        self.assertEqual(detect_script_type("12345"), "other")
        self.assertEqual(detect_script_type("---"), "other")
        self.assertEqual(detect_script_type(""), "other")

    def test_is_different_language_pair(self):
        # English + Chinese
        self.assertTrue(is_different_language_pair("Product Name", "产品名称"))
        self.assertTrue(is_different_language_pair("Section 1: Identification", "第一部分：化学品标识"))

        # English + Vietnamese
        self.assertTrue(is_different_language_pair("Product Name", "Tên sản phẩm"))

        # Chinese + Vietnamese
        self.assertTrue(is_different_language_pair("产品名称", "Tên sản phẩm"))

        # English + Russian
        self.assertTrue(is_different_language_pair("Identification", "Идентификация"))

        # Same language (English + English) -> Không phải song ngữ
        self.assertFalse(is_different_language_pair("Product Name", "Chemical Identifier"))
        self.assertFalse(is_different_language_pair("Wear eye protection", "Do not breathe dust"))

        # Trùng lặp hoàn toàn
        self.assertFalse(is_different_language_pair("Acetone", "Acetone"))

        # Chuỗi rỗng hoặc số thuần
        self.assertFalse(is_different_language_pair("", "产品名称"))
        self.assertFalse(is_different_language_pair("12345", "67890"))

        # Styled as translation override
        self.assertTrue(is_different_language_pair("Product Name", "Item Description", is_b_styled_as_trans=True))


class TestWordBilingualInPlace(unittest.TestCase):
    """Kiểm thử tính năng thay thế ngôn ngữ thứ 2 tại chỗ cho tệp Word (.docx)."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    @patch.object(InterlinearTranslator, "translate_texts_free")
    def test_word_in_paragraph_bilingual_replacement(self, mock_translate):
        # Mock dịch giả lập: dịch sang tiếng Việt
        mock_translate.side_effect = lambda texts, **kwargs: [f"Dịch: {t}" for t in texts]

        # 1. Tạo file docx có 1 đoạn văn chứa 2 dòng (Dòng 1: Tiếng Anh, Dòng 2: Tiếng Trung)
        doc = Document()
        p = doc.add_paragraph("Product Name\n产品名称")
        input_docx = os.path.join(self.temp_dir, "input_bilingual_multiline.docx")
        output_docx = os.path.join(self.temp_dir, "output_bilingual_multiline.docx")
        doc.save(input_docx)

        # 2. Thực hiện dịch
        res = InterlinearTranslator.translate_docx(
            input_path=input_docx,
            output_path=output_docx,
            source_lang="en",
            target_lang="vi",
            provider="free",
            output_mode="bilingual"
        )
        self.assertTrue(res["success"])

        # 3. Mở file kết quả kiểm tra
        res_doc = Document(output_docx)
        self.assertEqual(len(res_doc.paragraphs), 1)

        result_text = res_doc.paragraphs[0].text
        lines = result_text.splitlines()

        # Xác thực: Chỉ có đúng 2 dòng, KHÔNG bị nhân đôi thành 3 hoặc 4 dòng
        self.assertEqual(len(lines), 2, f"Số dòng phải là 2 nhưng thực tế là {len(lines)}: {lines}")
        self.assertEqual(lines[0], "Product Name")
        self.assertEqual(lines[1], "Dịch: Product Name")
        self.assertNotIn("产品名称", result_text, "Tiếng Trung cũ phải được thay thế hoàn toàn tại chỗ!")

        # Xác thực định dạng DOC-TRAN: dòng dịch mang màu xanh đậm #002060 và in nghiêng
        p_runs = res_doc.paragraphs[0].runs
        trans_runs = [r for r in p_runs if "Dịch: Product Name" in r.text]
        self.assertTrue(len(trans_runs) > 0)
        self.assertTrue(trans_runs[0].italic)
        self.assertEqual(trans_runs[0].font.color.rgb, DARK_BLUE)

    @patch.object(InterlinearTranslator, "translate_texts_free")
    def test_word_adjacent_paragraphs_bilingual_replacement(self, mock_translate):
        # Mock dịch giả lập
        mock_translate.side_effect = lambda texts, **kwargs: [f"Dịch: {t}" for t in texts]

        # 1. Tạo file docx có bảng (Table), trong ô có 2 đoạn văn liền kề (Anh + Trung)
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        # Đoạn 1: Tiếng Anh
        cell.paragraphs[0].text = "Composition / Information on Ingredients"
        # Đoạn 2: Tiếng Trung
        cell.add_paragraph("成分/组成信息")

        input_docx = os.path.join(self.temp_dir, "input_adjacent.docx")
        output_docx = os.path.join(self.temp_dir, "output_adjacent.docx")
        doc.save(input_docx)

        # 2. Dịch
        res = InterlinearTranslator.translate_docx(
            input_path=input_docx,
            output_path=output_docx,
            source_lang="en",
            target_lang="vi",
            provider="free",
            output_mode="bilingual"
        )
        self.assertTrue(res["success"])

        # 3. Kiểm tra file kết quả
        res_doc = Document(output_docx)
        res_cell = res_doc.tables[0].cell(0, 0)

        # Số đoạn văn trong ô vẫn là đúng 2 đoạn (không sinh thêm đoạn thứ 3 hoặc thứ 4)
        self.assertEqual(len(res_cell.paragraphs), 2)
        self.assertEqual(res_cell.paragraphs[0].text, "Composition / Information on Ingredients")
        self.assertEqual(res_cell.paragraphs[1].text, "Dịch: Composition / Information on Ingredients")
        self.assertNotIn("成分/组成信息", res_cell.text, "Tiếng Trung cũ ở đoạn 2 phải được thay thế tại chỗ!")

        # Kiểm tra màu sắc và kiểu dáng đoạn 2
        p2_runs = res_cell.paragraphs[1].runs
        self.assertTrue(len(p2_runs) > 0)
        self.assertTrue(p2_runs[0].italic)
        self.assertEqual(p2_runs[0].font.color.rgb, DARK_BLUE)

    @patch.object(InterlinearTranslator, "translate_texts_free")
    def test_word_monolingual_still_interlinear(self, mock_translate):
        # Đoạn văn đơn ngữ thông thường: Vẫn chèn dòng dịch xen kẽ bên dưới bình thường
        mock_translate.side_effect = lambda texts, **kwargs: [f"Dịch: {t}" for t in texts]

        doc = Document()
        doc.add_paragraph("First-aid measures")
        input_docx = os.path.join(self.temp_dir, "input_mono.docx")
        output_docx = os.path.join(self.temp_dir, "output_mono.docx")
        doc.save(input_docx)

        res = InterlinearTranslator.translate_docx(
            input_path=input_docx,
            output_path=output_docx,
            source_lang="en",
            target_lang="vi",
            provider="free",
            output_mode="bilingual"
        )
        self.assertTrue(res["success"])

        res_doc = Document(output_docx)
        lines = res_doc.paragraphs[0].text.splitlines()
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[0], "First-aid measures")
        self.assertEqual(lines[1], "Dịch: First-aid measures")


class TestExcelBilingualInPlace(unittest.TestCase):
    """Kiểm thử tính năng thay thế ngôn ngữ thứ 2 tại chỗ cho tệp Excel (.xlsx)."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    @patch.object(InterlinearTranslator, "translate_texts_free")
    def test_excel_bilingual_cell_replacement(self, mock_translate):
        mock_translate.side_effect = lambda texts, **kwargs: [f"Dịch: {t}" for t in texts]

        # 1. Tạo file Excel có 1 ô chứa song ngữ 2 dòng (Dòng 1: Tiếng Anh, Dòng 2: Tiếng Trung)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TestSheet"
        ws["A1"] = "Hazard Statement\n危险说明"
        ws["A2"] = "Single Line Monolingual"
        ws["A3"] = "=SUM(B1:B5)"  # Công thức toán học không được đụng đến

        input_xlsx = os.path.join(self.temp_dir, "input_bilingual.xlsx")
        output_xlsx = os.path.join(self.temp_dir, "output_bilingual.xlsx")
        wb.save(input_xlsx)

        # 2. Dịch bảng tính
        res = ExcelInterlinearTranslator.translate_excel(
            input_path=input_xlsx,
            output_path=output_xlsx,
            source_lang="en",
            target_lang="vi",
            provider="google",
            output_mode="bilingual"
        )
        self.assertTrue(res["success"])

        # 3. Mở file kết quả
        res_wb = openpyxl.load_workbook(output_xlsx, rich_text=True, data_only=False)
        res_ws = res_wb["TestSheet"]

        # Kiểm tra ô A1 (Song ngữ có sẵn)
        cell_a1_val = res_ws["A1"].value
        self.assertIsInstance(cell_a1_val, CellRichText)
        text_a1 = str(cell_a1_val)
        lines_a1 = text_a1.splitlines()

        # Phải có đúng 2 dòng: dòng 1 tiếng Anh, dòng 2 tiếng Việt dịch mới
        self.assertEqual(len(lines_a1), 2, f"Ô A1 phải có đúng 2 dòng nhưng thực tế là {len(lines_a1)}: {lines_a1}")
        self.assertEqual(lines_a1[0], "Hazard Statement")
        self.assertEqual(lines_a1[1], "Dịch: Hazard Statement")
        self.assertNotIn("危险说明", text_a1, "Dòng tiếng Trung cũ phải được thay thế hoàn toàn tại chỗ!")

        # Kiểm tra định dạng rich text của dòng dịch trong ô A1
        blocks = cell_a1_val
        self.assertGreaterEqual(len(blocks), 2)
        trans_block = blocks[1]
        self.assertTrue(trans_block.font.i, "Dòng dịch phải in nghiêng")
        self.assertEqual(trans_block.font.color.rgb, ExcelInterlinearTranslator.DARK_BLUE_ARGB, "Dòng dịch phải mang màu FF002060")

        # Kiểm tra ô A2 (Đơn ngữ): Vẫn được chèn xen kẽ bình thường
        cell_a2_val = res_ws["A2"].value
        self.assertIsInstance(cell_a2_val, CellRichText)
        lines_a2 = str(cell_a2_val).splitlines()
        self.assertEqual(len(lines_a2), 2)
        self.assertEqual(lines_a2[0], "Single Line Monolingual")
        self.assertEqual(lines_a2[1], "Dịch: Single Line Monolingual")

        # Kiểm tra ô A3 (Công thức): Giữ nguyên 100%
        self.assertEqual(res_ws["A3"].value, "=SUM(B1:B5)")

    @patch.object(InterlinearTranslator, "translate_texts_free")
    def test_excel_four_lines_two_pairs(self, mock_translate):
        mock_translate.side_effect = lambda texts, **kwargs: [f"Dịch: {t}" for t in texts]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["B2"] = "Product Name\n产品名称\nManufacturer\n生产企业"

        input_xlsx = os.path.join(self.temp_dir, "input_4lines.xlsx")
        output_xlsx = os.path.join(self.temp_dir, "output_4lines.xlsx")
        wb.save(input_xlsx)

        res = ExcelInterlinearTranslator.translate_excel(
            input_path=input_xlsx,
            output_path=output_xlsx,
            source_lang="en",
            target_lang="vi",
            provider="google",
            output_mode="bilingual"
        )
        self.assertTrue(res["success"])

        res_wb = openpyxl.load_workbook(output_xlsx, rich_text=True, data_only=False)
        cell_val = res_wb.active["B2"].value
        self.assertIsInstance(cell_val, CellRichText)
        lines = str(cell_val).splitlines()

        # Số dòng phải giữ nguyên đúng 4 dòng (2 cặp song ngữ)
        self.assertEqual(len(lines), 4)
        self.assertEqual(lines[0], "Product Name")
        self.assertEqual(lines[1], "Dịch: Product Name")
        self.assertEqual(lines[2], "Manufacturer")
        self.assertEqual(lines[3], "Dịch: Manufacturer")
        self.assertNotIn("产品名称", str(cell_val))
        self.assertNotIn("生产企业", str(cell_val))


class TestWordAdvancedPairs(unittest.TestCase):
    """Kiểm thử mở rộng cho Word: 4 dòng và thay thế tiếng Việt cũ."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    @patch.object(InterlinearTranslator, "translate_texts_free")
    def test_word_four_lines_two_pairs(self, mock_translate):
        mock_translate.side_effect = lambda texts, **kwargs: [f"Dịch: {t}" for t in texts]

        doc = Document()
        doc.add_paragraph("Product Name\n产品名称\nManufacturer\n生产企业")
        input_docx = os.path.join(self.temp_dir, "input_4lines.docx")
        output_docx = os.path.join(self.temp_dir, "output_4lines.docx")
        doc.save(input_docx)

        res = InterlinearTranslator.translate_docx(
            docx_path=input_docx,
            output_path=output_docx,
            source_lang="en",
            target_lang="vi",
            provider="free",
            output_mode="bilingual"
        )
        self.assertTrue(res["success"])

        res_doc = Document(output_docx)
        lines = res_doc.paragraphs[0].text.splitlines()
        self.assertEqual(len(lines), 4)
        self.assertEqual(lines[0], "Product Name")
        self.assertEqual(lines[1], "Dịch: Product Name")
        self.assertEqual(lines[2], "Manufacturer")
        self.assertEqual(lines[3], "Dịch: Manufacturer")
        self.assertNotIn("产品名称", res_doc.paragraphs[0].text)
        self.assertNotIn("生产企业", res_doc.paragraphs[0].text)

    @patch.object(InterlinearTranslator, "translate_texts_free")
    def test_word_english_vietnamese_old_replacement(self, mock_translate):
        mock_translate.side_effect = lambda texts, **kwargs: [f"Bản dịch mới: {t}" for t in texts]

        doc = Document()
        # Dòng 1: Tiếng Anh, Dòng 2: Tiếng Việt cũ
        doc.add_paragraph("Hazard Statement\nCác cảnh báo nguy hại cũ")
        input_docx = os.path.join(self.temp_dir, "input_vi_old.docx")
        output_docx = os.path.join(self.temp_dir, "output_vi_old.docx")
        doc.save(input_docx)

        res = InterlinearTranslator.translate_docx(
            docx_path=input_docx,
            output_path=output_docx,
            source_lang="en",
            target_lang="vi",
            provider="free",
            output_mode="bilingual"
        )
        self.assertTrue(res["success"])

        res_doc = Document(output_docx)
        lines = res_doc.paragraphs[0].text.splitlines()
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[0], "Hazard Statement")
        self.assertEqual(lines[1], "Bản dịch mới: Hazard Statement")
        self.assertNotIn("Các cảnh báo nguy hại cũ", res_doc.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
